"""Reproducible fixture builder for 08-nested-lists.

Two-level numbered + bullet hierarchy. python-docx applies list styles via
the built-in ``List Number`` / ``List Bullet`` style families; pandoc
recognises them and emits canonical markdown lists.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = Path(__file__).parent / "input.docx"


def build() -> None:
    doc = Document()
    doc.add_paragraph("Упорядоченный список:")

    doc.add_paragraph("Первый пункт", style="List Number")
    doc.add_paragraph("Подпункт первого", style="List Number 2")
    doc.add_paragraph("Ещё подпункт", style="List Number 2")
    doc.add_paragraph("Второй пункт", style="List Number")
    doc.add_paragraph("Подпункт второго", style="List Number 2")
    doc.add_paragraph("Третий пункт", style="List Number")

    doc.add_paragraph("Маркированный список:")

    doc.add_paragraph("Первое наблюдение", style="List Bullet")
    doc.add_paragraph("Дочернее наблюдение", style="List Bullet 2")
    doc.add_paragraph("Ещё одно дочернее", style="List Bullet 2")
    doc.add_paragraph("Второе наблюдение", style="List Bullet")

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(f"wrote {OUT}")
