"""Reproducible fixture builder for 09-pdf-text (T041).

Builds ``input.pdf`` by:

1. Constructing a deterministic DOCX with python-docx (heading + paragraphs).
2. Converting that DOCX → PDF via a live unoserver (LibreOffice).

Run inside the docker-compose stack where unoserver is reachable:

    poetry run python tests/import/golden/09-pdf-text/generate.py

The PDF is committed to the repo so CI does not need unoserver to *read*
the fixture — only to bootstrap a new one. To refresh the markdown baseline
after the pipeline changes, run with ``MARKDOWN_GOST_UPDATE_IMPORT_GOLDEN=1``.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from markdown_gost.output.pdf_writer import convert_to_pdf

HERE = Path(__file__).parent
DOCX_PATH = HERE / "_source.docx"
PDF_PATH = HERE / "input.pdf"


def _build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("PDF golden case", level=1)
    doc.add_paragraph(
        "Это первый абзац фикстуры. PDF собирается из этого docx через unoserver."
    )
    doc.add_paragraph(
        "Второй короткий абзац — нужен, чтобы документ не считался пустым."
    )
    doc.save(path)


def build() -> None:
    _build_docx(DOCX_PATH)
    pdf_bytes = convert_to_pdf(DOCX_PATH.read_bytes())
    PDF_PATH.write_bytes(pdf_bytes)
    DOCX_PATH.unlink(missing_ok=True)


if __name__ == "__main__":
    build()
    print(f"wrote {PDF_PATH} ({PDF_PATH.stat().st_size} bytes)")
