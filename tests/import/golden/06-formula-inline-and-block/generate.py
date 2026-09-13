"""Reproducible fixture builder for 06-formula-inline-and-block.

python-docx has no high-level OMML helper, so we inject raw
``<m:oMath>`` / ``<m:oMathPara>`` XML into the body via the element tree.
The simple formulas below (``a^2 + b^2 = c^2``, ``E = mc^2``) are enough
to exercise pandoc's OMML→TeX bridge.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from lxml import etree

OUT = Path(__file__).parent / "input.docx"

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NSMAP = {"m": M_NS}


def _math_text(value: str) -> etree._Element:
    r = etree.SubElement(etree.Element("placeholder"), f"{{{M_NS}}}r")
    t = etree.SubElement(r, f"{{{M_NS}}}t")
    t.text = value
    return r


def _build_quadratic() -> etree._Element:
    """Build OMML for ``a^2 + b^2 = c^2`` as a single inline ``oMath``."""
    omath = etree.Element(f"{{{M_NS}}}oMath", nsmap=NSMAP)

    def superscript(base: str, exp: str) -> etree._Element:
        sup = etree.SubElement(omath, f"{{{M_NS}}}sSup")
        e = etree.SubElement(sup, f"{{{M_NS}}}e")
        e.append(_math_text(base))
        sup_e = etree.SubElement(sup, f"{{{M_NS}}}sup")
        sup_e.append(_math_text(exp))
        return sup

    superscript("a", "2")
    omath.append(_math_text(" + "))
    superscript("b", "2")
    omath.append(_math_text(" = "))
    superscript("c", "2")
    return omath


def _build_einstein() -> etree._Element:
    """``E = mc^2`` as block math (``oMathPara > oMath``)."""
    omath_para = etree.Element(f"{{{M_NS}}}oMathPara", nsmap=NSMAP)
    omath = etree.SubElement(omath_para, f"{{{M_NS}}}oMath")
    omath.append(_math_text("E = m"))
    sup = etree.SubElement(omath, f"{{{M_NS}}}sSup")
    e = etree.SubElement(sup, f"{{{M_NS}}}e")
    e.append(_math_text("c"))
    sup_e = etree.SubElement(sup, f"{{{M_NS}}}sup")
    sup_e.append(_math_text("2"))
    return omath_para


def build() -> None:
    doc = Document()

    inline_para = doc.add_paragraph("Известная формула ")
    inline_para._p.append(_build_quadratic())
    inline_para.add_run(" — теорема Пифагора.")

    doc.add_paragraph("Уравнение массы-энергии Эйнштейна:")
    block_para = doc.add_paragraph()
    block_para._p.append(_build_einstein())

    doc.add_paragraph("После формул — обычный абзац.")

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(f"wrote {OUT}")
