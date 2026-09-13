"""Reproducible fixture builder for 04-image-with-caption."""

from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.shared import Cm

OUT = Path(__file__).parent / "input.docx"


def _png_bytes() -> bytes:
    sig = b"\x89PNG\r\n\x1a\n"

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0)
    raw = b"\x00\xff\x00\x00\x00\x00\xff\x00\x00\xff\x00\x00\x00\x00\x00"
    idat = zlib.compress(raw)
    return sig + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


def build() -> None:
    doc = Document()
    doc.add_paragraph("Перед иллюстрацией.")
    doc.add_picture(io.BytesIO(_png_bytes()), width=Cm(2))
    doc.add_paragraph("Рисунок 1 — Схема архитектуры")
    doc.add_paragraph("После иллюстрации.")
    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(f"wrote {OUT}")
