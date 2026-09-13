"""Reproducible fixture builder for 01-simple-text.

Run from the repo root:

    poetry run python tests/import/golden/01-simple-text/generate.py

Produces ``input.docx`` next to this script.  Check it in alongside the
expected.md baseline; CI only runs the harness, not the generator.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = Path(__file__).parent / "input.docx"


def build() -> None:
    doc = Document()

    doc.add_paragraph("Это первый абзац обычного текста без форматирования.")

    second = doc.add_paragraph()
    second.add_run("Во втором абзаце есть ")
    second.add_run("жирный").bold = True
    second.add_run(", ")
    second.add_run("курсив").italic = True
    second.add_run(" и ")
    second.add_run("подчёркивание").underline = True
    second.add_run(".")

    third = doc.add_paragraph()
    third.add_run("Здесь — ")
    strike = third.add_run("зачёркнутый")
    strike.font.strike = True
    third.add_run(" фрагмент.")

    doc.add_paragraph("Финальный короткий абзац.")

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(f"wrote {OUT}")
