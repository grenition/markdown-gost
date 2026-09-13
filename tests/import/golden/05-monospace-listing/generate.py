"""Reproducible fixture builder for 05-monospace-listing.

Monospace-fonted paragraphs trigger the listing detector (T035); a
``Листинг N — ...`` caption above them must become a trailing ``: Caption``.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = Path(__file__).parent / "input.docx"

CODE_LINES = (
    "def greet(name):",
    "    print(f\"Hello, {name}!\")",
    "",
    "greet(\"world\")",
)


def _add_monospace_paragraph(doc, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Courier New"


def build() -> None:
    doc = Document()
    doc.add_paragraph("Ниже приведён пример программы.")
    doc.add_paragraph("Листинг 1 — Приветствие пользователя")
    for line in CODE_LINES:
        _add_monospace_paragraph(doc, line)
    doc.add_paragraph("Конец примера.")
    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(f"wrote {OUT}")
