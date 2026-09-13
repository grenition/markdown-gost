"""Reproducible fixture builder for 03-table-with-caption."""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = Path(__file__).parent / "input.docx"


def build() -> None:
    doc = Document()

    doc.add_paragraph("Перед таблицей идёт обычный абзац.")
    doc.add_paragraph("Таблица 1 — Сравнение характеристик")

    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "Параметр"
    headers[1].text = "Значение A"
    headers[2].text = "Значение B"

    rows_data = [
        ("Скорость", "10 МБ/с", "25 МБ/с"),
        ("Цена", "1000 ₽", "1500 ₽"),
    ]
    for row_idx, row_data in enumerate(rows_data, start=1):
        cells = table.rows[row_idx].cells
        for col_idx, value in enumerate(row_data):
            cells[col_idx].text = value

    doc.add_paragraph("После таблицы — снова обычный абзац.")

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(f"wrote {OUT}")
