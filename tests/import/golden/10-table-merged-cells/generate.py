"""Reproducible fixture builder for 10-table-merged-cells.

Builds a DOCX whose table has a horizontally merged cell. Pandoc cannot
express such a layout in ``pipe_tables`` and falls back to raw HTML —
this case exercises the T046 html-table converter end-to-end.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = Path(__file__).parent / "input.docx"


def build() -> None:
    doc = Document()

    doc.add_paragraph("Перед таблицей идёт обычный абзац.")

    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"

    headers = table.rows[0].cells
    headers[0].text = "A"
    headers[1].text = "B"
    headers[2].text = "C"

    # Row 1: regular data.
    row1 = table.rows[1].cells
    row1[0].text = "x"
    row1[1].text = "y"
    row1[2].text = "z"

    # Row 2: horizontally merged cell across all three columns.
    row2 = table.rows[2].cells
    merged = row2[0].merge(row2[1]).merge(row2[2])
    merged.text = "Объединённый текст"

    doc.add_paragraph("После таблицы — обычный абзац.")

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(f"wrote {OUT}")
