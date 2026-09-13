from __future__ import annotations

from docx import Document
from docx.shared import Cm, Length, Pt

from markdown_gost.render.renderer import Renderer
from markdown_gost.renderable.base import Renderable, RenderedInfo


class _FixedHeightRenderable(Renderable):
    """Test double: yields exactly one RenderedInfo with a fixed height."""

    def __init__(self, height: Length, document: Document) -> None:
        self._height = height
        # We need a real docx element to append; use a fresh paragraph.
        self._document = document

    def render(self, previous_rendered, layout_state):  # type: ignore[no-untyped-def]
        paragraph = self._document.add_paragraph("")
        # remove the just-added paragraph from the body so the renderer can
        # re-append it itself (the renderer is responsible for placement)
        paragraph._p.getparent().remove(paragraph._p)
        yield RenderedInfo(docx_element=paragraph, height=self._height)


def _document_with_known_page_size() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    return document


def test_empty_renderables_leaves_document_empty() -> None:
    document = _document_with_known_page_size()
    body_paragraphs_before = len(document.paragraphs)

    renderer = Renderer(document)
    renderer.process([])

    # No new paragraphs should appear except whatever the renderer already
    # touched (footer page-number field is in the footer, not the body).
    assert len(document.paragraphs) == body_paragraphs_before


def test_single_renderable_advances_layout_height() -> None:
    document = _document_with_known_page_size()
    renderer = Renderer(document)

    height = Pt(72)  # ~1 inch
    renderer.process([_FixedHeightRenderable(height, document)])

    assert renderer.layout_tracker.current_state.current_page_height == height
    assert renderer.layout_tracker.current_state.page == 1


def test_oversized_renderable_is_pushed_to_next_page() -> None:
    document = _document_with_known_page_size()
    renderer = Renderer(document)

    # First, fill most of page 1.
    fill = renderer.layout_tracker.current_state.remaining_page_height - Pt(10)
    renderer.process([_FixedHeightRenderable(Length(fill), document)])
    assert renderer.layout_tracker.current_state.page == 1

    # Now render something taller than the remaining page space (~10pt left).
    big = Pt(200)
    renderer.process([_FixedHeightRenderable(big, document)])

    # Renderer should have flushed to a new page before placing the big item.
    assert renderer.layout_tracker.current_state.page == 2
