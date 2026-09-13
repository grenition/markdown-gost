"""End-to-end CLI import test (T036).

Spawns ``python -m markdown_gost.cli import ...`` against a real DOCX produced by
``python-docx`` and verifies the resulting markdown plus side-effect files
(images directory) live where the contract promises.
"""

from __future__ import annotations

import hashlib
import os
import shutil
import struct
import subprocess
import sys
import zlib
from pathlib import Path

import pytest
from docx import Document

from markdown_gost.import_.pandoc_runner import DEFAULT_PANDOC_BINARY

pytestmark = pytest.mark.requires_pandoc


def _pandoc_available() -> bool:
    binary = os.environ.get("PANDOC_BINARY", DEFAULT_PANDOC_BINARY)
    return shutil.which(binary) is not None


@pytest.fixture(autouse=True)
def _skip_if_no_pandoc() -> None:
    if not _pandoc_available():
        pytest.skip("pandoc is not installed; run inside docker image")


def _png_bytes() -> bytes:
    sig = b"\x89PNG\r\n\x1a\n"

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = bytes([0, 255, 0, 0])
    idat = zlib.compress(raw)
    return sig + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


def _build_docx(path: Path) -> None:
    import io

    doc = Document()
    doc.add_heading("Заголовок", level=1)
    doc.add_paragraph("Первый абзац.")
    doc.add_heading("Подзаголовок", level=2)
    doc.add_picture(io.BytesIO(_png_bytes()))
    doc.add_paragraph("Второй абзац.")
    doc.save(path)


def _run_cli(*args: str, cwd: Path) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "markdown_gost.cli", *args],
        capture_output=True,
        text=True,
        cwd=cwd,
    )


def test_import_real_docx_writes_md_and_images(tmp_path: Path) -> None:
    docx = tmp_path / "doc.docx"
    _build_docx(docx)

    result = _run_cli("import", str(docx), cwd=tmp_path)

    assert result.returncode == 0, result.stderr
    out_md = tmp_path / "doc.md"
    assert out_md.exists()
    body = out_md.read_text(encoding="utf-8")
    assert "Заголовок" in body
    assert "Подзаголовок" in body
    assert "Первый абзац" in body
    # default images dir = <output-stem>_files/
    images_dir = tmp_path / "doc_files"
    assert images_dir.is_dir()
    image_id = hashlib.sha256(_png_bytes()).hexdigest()[:32]
    written = sorted(p.name for p in images_dir.iterdir())
    assert written == [image_id]
    assert f"doc_files/{image_id}" in body
    assert "Imported:" in result.stderr


def test_import_explicit_output_and_images_dir(tmp_path: Path) -> None:
    docx = tmp_path / "src.docx"
    _build_docx(docx)
    out_md = tmp_path / "out" / "result.md"
    images = tmp_path / "media"

    result = _run_cli(
        "import",
        str(docx),
        "-o",
        str(out_md),
        "--images-dir",
        str(images),
        cwd=tmp_path,
    )

    assert result.returncode == 0, result.stderr
    assert out_md.exists()
    assert images.is_dir()
    image_id = hashlib.sha256(_png_bytes()).hexdigest()[:32]
    assert (images / image_id).exists()


def test_import_unsupported_extension_returns_exit_1(tmp_path: Path) -> None:
    src = tmp_path / "in.txt"
    src.write_bytes(b"hello")
    result = _run_cli("import", str(src), cwd=tmp_path)
    assert result.returncode == 1
    assert "unsupported" in result.stderr.lower()


def test_import_missing_input_returns_exit_1(tmp_path: Path) -> None:
    result = _run_cli("import", str(tmp_path / "nope.docx"), cwd=tmp_path)
    assert result.returncode == 1
