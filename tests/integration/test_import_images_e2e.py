"""End-to-end import of a DOCX with images (T034).

Builds a real docx via ``python-docx`` with two embedded images, runs pandoc,
and checks both the local-FS and S3 routing paths.

* Pandoc has to be on PATH (or via ``PANDOC_BINARY``) — uses the
  ``requires_pandoc`` marker.
* The S3 path is exercised only when ``S3_ENDPOINT`` / ``S3_ACCESS_KEY`` /
  ``S3_SECRET_KEY`` are set (MinIO from docker-compose); otherwise skipped.
"""

from __future__ import annotations

import contextlib
import hashlib
import io
import os
import shutil
import struct
import uuid
import zlib
from pathlib import Path

import pytest
from docx import Document

from markdown_gost.import_ import ImportContext, import_docx
from markdown_gost.import_.pandoc_runner import DEFAULT_PANDOC_BINARY
from markdown_gost.storage import FilesystemStorage
from markdown_gost.storage.s3 import S3Storage

pytestmark = pytest.mark.requires_pandoc


def _content_id(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:32]


def _pandoc_available() -> bool:
    binary = os.environ.get("PANDOC_BINARY", DEFAULT_PANDOC_BINARY)
    return shutil.which(binary) is not None


@pytest.fixture(autouse=True)
def _skip_if_no_pandoc() -> None:
    if not _pandoc_available():
        pytest.skip("pandoc is not installed; run inside docker image")


def _png_bytes(rgb: tuple[int, int, int]) -> bytes:
    """Minimal valid 1×1 PNG with the given RGB pixel."""
    sig = b"\x89PNG\r\n\x1a\n"

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = bytes([0, rgb[0], rgb[1], rgb[2]])
    idat = zlib.compress(raw)
    return sig + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


def _build_docx_with_two_images(path: Path) -> None:
    doc = Document()
    doc.add_heading("Two images", level=1)
    doc.add_paragraph("Before first.")
    doc.add_picture(io.BytesIO(_png_bytes((255, 0, 0))))
    doc.add_paragraph("Between images.")
    doc.add_picture(io.BytesIO(_png_bytes((0, 255, 0))))
    doc.add_paragraph("After second.")
    doc.save(path)


def test_local_handler_writes_two_images_and_rewrites_md(tmp_path: Path) -> None:
    docx_path = tmp_path / "in.docx"
    _build_docx_with_two_images(docx_path)

    out_dir = tmp_path / "result_files"
    ctx = ImportContext(
        storage=FilesystemStorage(base_dir=tmp_path),
        images_prefix=None,
        images_dir=out_dir,
    )

    result = import_docx(docx_path, ctx)

    expected_ids = {
        _content_id(_png_bytes((255, 0, 0))),
        _content_id(_png_bytes((0, 255, 0))),
    }
    written = {p.name for p in out_dir.iterdir()}
    assert written == expected_ids
    # No file has an extension; every name is a 32-hex id.
    for p in out_dir.iterdir():
        assert p.suffix == ""
        assert len(p.name) == 32
    for image_id in expected_ids:
        assert f"result_files/{image_id}" in result.markdown
    assert {img.name for img in result.images} == expected_ids
    assert result.fallbacks.get("image", 0) == 0


def _env_or_none(name: str) -> str | None:
    value = os.environ.get(name)
    return value if value else None


@pytest.fixture
def s3_client_and_bucket():
    endpoint = _env_or_none("S3_ENDPOINT")
    access = _env_or_none("S3_ACCESS_KEY")
    secret = _env_or_none("S3_SECRET_KEY")
    if not (endpoint and access and secret):
        pytest.skip("S3_* env not set; skipping MinIO part")
    bucket = os.environ.get("S3_BUCKET", "markdown-gost")
    region = os.environ.get("S3_REGION") or "us-east-1"
    boto3 = pytest.importorskip("boto3")
    client = boto3.client(
        "s3",
        endpoint_url=endpoint,
        aws_access_key_id=access,
        aws_secret_access_key=secret,
        region_name=region,
    )
    try:
        client.head_bucket(Bucket=bucket)
    except Exception:
        try:
            client.create_bucket(Bucket=bucket)
        except Exception as exc:
            pytest.skip(f"cannot prepare bucket {bucket!r}: {exc}")
    return client, bucket


def test_storage_handler_uploads_two_images_to_minio(
    tmp_path: Path, s3_client_and_bucket
) -> None:
    client, bucket = s3_client_and_bucket
    docx_path = tmp_path / "in.docx"
    _build_docx_with_two_images(docx_path)

    prefix = f"md2gost-test/{uuid.uuid4().hex}"
    storage = S3Storage(bucket=bucket, client=client)
    ctx = ImportContext(
        storage=storage,
        images_prefix=prefix,
        images_dir=None,
    )

    expected_ids = (
        _content_id(_png_bytes((255, 0, 0))),
        _content_id(_png_bytes((0, 255, 0))),
    )
    expected_keys = [f"{prefix}/{image_id}" for image_id in expected_ids]
    try:
        result = import_docx(docx_path, ctx)
        for key in expected_keys:
            assert key in result.markdown
            client.head_object(Bucket=bucket, Key=key)
        assert {img.name for img in result.images} == set(expected_ids)
    finally:
        for key in expected_keys:
            with contextlib.suppress(Exception):
                client.delete_object(Bucket=bucket, Key=key)
