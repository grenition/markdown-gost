from __future__ import annotations

from io import BytesIO

from docx import Document

from markdown_gost.config.loader import load_config_from_string
from markdown_gost.convert import convert


def test_convert_docx_renders_citations_and_bibliography_order() -> None:
    cfg = load_config_from_string("preset: default\n")
    markdown = (
        "# ВВЕДЕНИЕ\n\n"
        "Первая ссылка [@b] и повтор [@b].\n\n"
        "- Пункт со второй ссылкой [@a]\n\n"
        "| Колонка |\n"
        "|---------|\n"
        "| Табличная ссылка [@c] |\n\n"
        "# СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ {.unnumbered}\n\n"
        "::: {.bibliography}\n"
        "- id: a\n"
        "  text: Автор А. Источник А.\n"
        "- id: b\n"
        "  text: Автор Б. Источник Б.\n"
        "- id: c\n"
        "  text: Автор В. Источник В.\n"
        ":::\n"
    )

    data = convert(markdown, cfg, format="docx")
    document = Document(BytesIO(data))

    paragraph_texts = [p.text for p in document.paragraphs if p.text]
    table_texts = [
        p.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for p in cell.paragraphs
    ]
    all_text = "\n".join([*paragraph_texts, *table_texts])

    assert "Первая ссылка [1] и повтор [1]." in all_text
    assert "Пункт со второй ссылкой [2]" in all_text
    assert "Табличная ссылка [3]" in all_text
    assert "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ" in all_text

    entries = [text for text in paragraph_texts if text[:3] in {"1. ", "2. ", "3. "}]
    assert entries == [
        "1. Автор Б. Источник Б.",
        "2. Автор А. Источник А.",
        "3. Автор В. Источник В.",
    ]
