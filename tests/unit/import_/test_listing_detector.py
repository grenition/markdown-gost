"""Unit tests for the monospace-paragraph detector (T035).

We build the fixture docx in-memory with ``python-docx`` so the tests do not
carry binary blobs in the repo.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocxDocument

from markdown_gost.import_.listing_detector import (
    ParagraphRange,
    scan_monospace_paragraphs,
)


def _add_run(paragraph, text: str, font_name: str | None) -> None:
    run = paragraph.add_run(text)
    if font_name is not None:
        run.font.name = font_name


def _save(doc: DocxDocument, tmp_path: Path, name: str = "fixture.docx") -> Path:
    path = tmp_path / name
    doc.save(path)
    return path


def test_returns_empty_for_document_without_monospace(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    _add_run(p, "Plain prose paragraph.", "Times New Roman")
    path = _save(doc, tmp_path)

    assert scan_monospace_paragraphs(path) == []


def test_detects_single_monospace_paragraph(tmp_path: Path) -> None:
    doc = Document()
    _add_run(doc.add_paragraph(), "Intro prose.", "Times New Roman")
    _add_run(doc.add_paragraph(), "print('hello')", "Courier New")
    _add_run(doc.add_paragraph(), "Outro prose.", "Times New Roman")
    path = _save(doc, tmp_path)

    ranges = scan_monospace_paragraphs(path)
    assert ranges == [ParagraphRange(start_index=1, texts=("print('hello')",))]


def test_groups_consecutive_monospace_paragraphs(tmp_path: Path) -> None:
    doc = Document()
    _add_run(doc.add_paragraph(), "Intro.", "Times New Roman")
    _add_run(doc.add_paragraph(), "def f():", "Consolas")
    _add_run(doc.add_paragraph(), "    return 1", "Consolas")
    _add_run(doc.add_paragraph(), "    return 2", "Monaco")
    _add_run(doc.add_paragraph(), "Outro.", "Times New Roman")
    path = _save(doc, tmp_path)

    ranges = scan_monospace_paragraphs(path)
    assert ranges == [
        ParagraphRange(
            start_index=1,
            texts=("def f():", "    return 1", "    return 2"),
        )
    ]


def test_breaks_range_on_non_monospace_paragraph(tmp_path: Path) -> None:
    doc = Document()
    _add_run(doc.add_paragraph(), "A = 1", "Courier New")
    _add_run(doc.add_paragraph(), "interlude", "Times New Roman")
    _add_run(doc.add_paragraph(), "B = 2", "Courier New")
    path = _save(doc, tmp_path)

    ranges = scan_monospace_paragraphs(path)
    assert ranges == [
        ParagraphRange(start_index=0, texts=("A = 1",)),
        ParagraphRange(start_index=2, texts=("B = 2",)),
    ]


def test_breaks_range_on_empty_paragraph(tmp_path: Path) -> None:
    doc = Document()
    _add_run(doc.add_paragraph(), "line1", "Courier New")
    doc.add_paragraph()  # truly empty paragraph
    _add_run(doc.add_paragraph(), "line2", "Courier New")
    path = _save(doc, tmp_path)

    ranges = scan_monospace_paragraphs(path)
    assert ranges == [
        ParagraphRange(start_index=0, texts=("line1",)),
        ParagraphRange(start_index=2, texts=("line2",)),
    ]


def test_paragraph_with_mixed_fonts_is_not_monospace(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    _add_run(p, "code = ", "Courier New")
    _add_run(p, "value", "Times New Roman")
    path = _save(doc, tmp_path)

    assert scan_monospace_paragraphs(path) == []


def test_ignores_empty_runs_when_classifying(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    _add_run(p, "", "Times New Roman")  # empty run with non-mono font, ignored
    _add_run(p, "x = 1", "Consolas")
    path = _save(doc, tmp_path)

    ranges = scan_monospace_paragraphs(path)
    assert ranges == [ParagraphRange(start_index=0, texts=("x = 1",))]


def test_recognises_all_listed_monospace_fonts(tmp_path: Path) -> None:
    fonts = (
        "Courier New",
        "Consolas",
        "Monaco",
        "Menlo",
        "Source Code Pro",
        "Liberation Mono",
        "DejaVu Sans Mono",
    )
    for font in fonts:
        doc = Document()
        _add_run(doc.add_paragraph(), "code", font)
        path = _save(doc, tmp_path, f"{font.replace(' ', '_')}.docx")
        assert scan_monospace_paragraphs(path) != [], font


def test_paragraph_without_font_name_is_not_monospace(tmp_path: Path) -> None:
    doc = Document()
    _add_run(doc.add_paragraph(), "no font set", None)
    path = _save(doc, tmp_path)

    assert scan_monospace_paragraphs(path) == []


def test_raises_for_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        scan_monospace_paragraphs(tmp_path / "nope.docx")
