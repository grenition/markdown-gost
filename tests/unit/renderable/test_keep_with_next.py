"""w:keepNext связывает подпись с объектом при разрыве страницы.

Кейс 15d-gost-captions: подпись не должна оставаться одна внизу страницы.
- Подпись таблицы (в т.ч. «Продолжение таблицы N») — keepNext на подписи.
- Подпись картинки снизу — keepNext на параграфе самой картинки.
- Шапка таблицы — keepNext на параграфах ячеек шапки (шапка не остаётся
  одна внизу страницы без первой строки тела).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx.shared import Cm
from docx.table import Table as DocxTable

from markdown_gost.config.loader import load_config_from_string
from markdown_gost.core.ast import nodes as ast
from markdown_gost.render.document_factory import build_document
from markdown_gost.render.layout_tracker import LayoutState
from markdown_gost.renderable.base import RenderedInfo
from markdown_gost.renderable.caption import Caption
from markdown_gost.renderable.image import Image
from markdown_gost.renderable.paragraph import Paragraph
from markdown_gost.renderable.table import Table
from markdown_gost.storage.fs import FilesystemStorage

FIXTURES = Path(__file__).parent / "_fixtures"


@pytest.fixture
def config():
    return load_config_from_string("preset: gost-7-32-2017\n")


@pytest.fixture
def document(config):
    return build_document(config)


@pytest.fixture
def storage():
    return FilesystemStorage(base_dir=FIXTURES)


@pytest.fixture
def layout(config):
    return LayoutState(max_height=Cm(20), max_width=Cm(15))


def test_table_caption_keeps_with_next(document, config):
    caption = Caption(
        document, config, category="table", number=2, text="Реестр", before=True
    )
    assert caption.docx_paragraph.paragraph_format.keep_with_next is True


def test_continuation_caption_keeps_with_next(document, config, layout):
    table = Table(document, config, ast.Table(rows=[]), caption_text="Реестр")
    para, _height = table._build_continuation(layout)
    assert para.paragraph_format.keep_with_next is True


def test_image_paragraph_keeps_with_caption(document, config, storage, layout):
    image = Image(
        document, config, ast.Image(src="sample.png", alt="Схема"), storage
    )
    image.set_number(1)
    items = list(image.render(None, layout))
    assert image.docx_paragraph.paragraph_format.keep_with_next is True
    # У подписи снизу keepNext нет — она последняя в связке.
    caption_para = items[-1].docx_element
    assert caption_para.paragraph_format.keep_with_next is None


def test_image_caption_below_has_no_keep_next(document, config):
    caption = Caption(
        document, config, category="image", number=1, text="Схема", before=False
    )
    assert caption.docx_paragraph.paragraph_format.keep_with_next is None


def test_ordinary_paragraph_has_no_keep_next(document, config, storage):
    paragraph = Paragraph(document, config, storage=storage)
    paragraph.add_inline_nodes([ast.Text(text="Обычный абзац")])
    assert paragraph.docx_paragraph.paragraph_format.keep_with_next is None


def _md_table() -> ast.Table:
    """Шапка + две body-строки на два столбца."""

    def row(header: bool, *texts: str) -> ast.TableRow:
        return ast.TableRow(
            header=header,
            cells=[ast.TableCell(children=[ast.Text(text=t)]) for t in texts],
        )

    return ast.Table(
        rows=[
            row(True, "Наименование", "Значение"),
            row(False, "Модуль питания", "12 В"),
            row(False, "Модуль обработки", "5 В"),
        ]
    )


def test_table_header_cells_keep_with_next(document, config):
    table = Table(document, config, _md_table(), caption_text="Реестр")
    assert table.docx_table is not None
    rows = table.docx_table.rows
    for cell in rows[0].cells:  # шапка → keepNext
        for para in cell.paragraphs:
            assert para.paragraph_format.keep_with_next is True
    for row in rows[1:]:  # body-строки → без keepNext
        for cell in row.cells:
            for para in cell.paragraphs:
                assert para.paragraph_format.keep_with_next is None


def test_continuation_table_header_cells_keep_with_next(config):
    config.captions.continuation_break = True
    document = build_document(config)
    # Низкая страница — таблица режется на несколько чанков с «Продолжение».
    layout = LayoutState(max_height=Cm(1), max_width=Cm(15))
    table = Table(document, config, _md_table(), caption_text="Реестр")
    table.set_number(2)
    chunks = [
        item.docx_element
        for item in table.render(None, layout)
        if isinstance(item, RenderedInfo) and isinstance(item.docx_element, DocxTable)
    ]
    assert len(chunks) >= 2  # разрез реально случился
    for tbl in chunks:
        for cell in tbl.rows[0].cells:  # шапка каждого чанка → keepNext
            for para in cell.paragraphs:
                assert para.paragraph_format.keep_with_next is True
        for row in tbl.rows[1:]:  # body-строки → без keepNext
            for cell in row.cells:
                for para in cell.paragraphs:
                    assert para.paragraph_format.keep_with_next is None
