"""Юнит-тесты шаблона ``content`` (TOC) — порт легаси-``ToC`` (T023).

Проверяем:
* шаблон рендерит один пустой параграф + PageBreak (height=0 у TOC, page-break
  выбирает остаток страницы),
* ``post_process`` создаёт отдельные параграфы записей с номером из
  ``RenderIndex`` и tab+page,
* ненумерованные заголовки попадают в TOC без префикса,
* schema strict: неизвестное поле → плейсхолдер.
"""

from __future__ import annotations

import pytest
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Length
from docx.text.paragraph import Paragraph as DocxParagraph

from markdown_gost.config.loader import load_config_from_string
from markdown_gost.render.document_factory import build_document
from markdown_gost.render.layout_tracker import LayoutTracker
from markdown_gost.render.render_index import RenderIndex
from markdown_gost.renderable.base import RenderedInfo
from markdown_gost.templates import TemplateContext, render_template
from markdown_gost.templates.content.render import Content


def _all_texts(p):
    """Все ``<w:t>`` параграфа — включая обёрнутые в ``<w:hyperlink>``.

    ``paragraph.runs`` пропускает runs внутри hyperlink-элемента, поэтому
    собираем тексты по XPath напрямую.
    """

    return [t.text or "" for t in p._p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")]


def _attach_toc_paragraph(rendered, document):
    document._body._element.append(rendered._docx_paragraph._element)


def _toc_paragraphs(rendered):
    first = rendered._docx_paragraph._element
    body = first.getparent()
    assert body is not None
    paragraphs = []
    current = first
    while current is not None and current.tag.endswith("}p"):
        paragraphs.append(DocxParagraph(current, rendered._parent))
        current = current.getnext()
    return paragraphs


def _tab_stops(paragraph):
    return list(paragraph.paragraph_format.tab_stops)


@pytest.fixture
def config():
    return load_config_from_string("preset: gost-7-32-2017\n")


@pytest.fixture
def document(config):
    return build_document(config)


@pytest.fixture
def ctx(document, config):
    return TemplateContext(document=document, config=config)


def test_content_default_params(ctx):
    rendered = render_template("content", {}, ctx)
    assert isinstance(rendered, Content)


def test_content_render_yields_paragraph_then_page_break(ctx):
    """С пустым ``title`` — TOC paragraph (height=0) + page-break."""

    rendered = render_template("content", {"title": ""}, ctx)
    tracker = LayoutTracker(Length(10_000_000), Length(6_000_000))
    infos = list(rendered.render(None, tracker.current_state))
    rendered_infos = [i for i in infos if isinstance(i, RenderedInfo)]
    assert len(rendered_infos) >= 2
    assert rendered_infos[0].height == 0
    last_xml = rendered_infos[-1].docx_element._p.xml
    assert 'w:type="page"' in last_xml


def test_content_unknown_param_rejected(document, config):
    rendered = render_template(
        "content",
        {"unknown": "value"},
        TemplateContext(document=document, config=config),
    )
    assert not isinstance(rendered, Content)


def test_content_post_process_writes_numbered_entries(ctx, document):
    rendered = render_template("content", {}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 0  # допускаем все заголовки в TOC (page > 0)
    _attach_toc_paragraph(rendered, document)

    index = RenderIndex()
    index.add_heading(level=1, text="Введение", numbered=True, number="7", page=1)
    index.add_heading(level=2, text="Цели", numbered=True, number="7.4", page=1)
    index.add_heading(level=2, text="Задачи", numbered=True, number="7.5", page=2)
    index.add_heading(level=1, text="Заключение", numbered=True, number="8", page=5)

    rendered.post_process(index, document)

    paragraphs = _toc_paragraphs(rendered)
    assert len(paragraphs) == 4

    expected = [
        ("7 ", "Введение", "1"),
        ("7.4 ", "Цели", "1"),
        ("7.5 ", "Задачи", "2"),
        ("8 ", "Заключение", "5"),
    ]
    for paragraph, (number, title, page) in zip(paragraphs, expected, strict=True):
        texts = _all_texts(paragraph)
        assert number in texts
        assert title in texts
        assert page in texts
        assert "<w:tab" in paragraph._p.xml
        assert "<w:br" not in paragraph._p.xml

    all_texts = [text for paragraph in paragraphs for text in _all_texts(paragraph)]
    assert "1. " not in all_texts
    assert "1.1. " not in all_texts


def test_content_unnumbered_heading_skips_prefix(ctx, document):
    rendered = render_template("content", {}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 0
    _attach_toc_paragraph(rendered, document)

    index = RenderIndex()
    index.add_heading(level=1, text="Введение", numbered=True, number="1", page=1)
    index.add_heading(
        level=1, text="Список литературы", numbered=False, number=None, page=10
    )
    rendered.post_process(index, document)

    paragraphs = _toc_paragraphs(rendered)
    texts = [text for paragraph in paragraphs for text in _all_texts(paragraph)]
    assert "1 " in texts
    assert "Введение" in texts
    assert "Список литературы" in texts
    # Префикс "2 " не должен появиться перед ненумерованным заголовком.
    assert "2 " not in texts


def test_content_filters_by_depth(ctx, document):
    rendered = render_template("content", {"depth": 1}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 0
    _attach_toc_paragraph(rendered, document)

    index = RenderIndex()
    index.add_heading(level=1, text="A", numbered=True, number="1", page=1)
    index.add_heading(level=2, text="B", numbered=True, number="1.1", page=1)
    index.add_heading(level=1, text="C", numbered=True, number="2", page=3)
    rendered.post_process(index, document)

    paragraphs = _toc_paragraphs(rendered)
    assert len(paragraphs) == 2
    texts = [text for paragraph in paragraphs for text in _all_texts(paragraph)]
    assert "A" in texts
    assert "C" in texts
    # B (level=2) выкинут, поэтому "1.1" не появляется и C берёт номер
    # напрямую из RenderIndex.
    assert "B" not in texts
    assert "2 " in texts
    assert "1.1 " not in texts


def test_content_includes_appendix_entries_and_prefixed_headings(ctx, document):
    rendered = render_template("content", {}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 0
    _attach_toc_paragraph(rendered, document)

    index = RenderIndex()
    index.add_heading(
        level=1,
        text="ПРИЛОЖЕНИЕ А Исходные данные",
        numbered=False,
        number=None,
        page=10,
    )
    index.add_heading(
        level=1,
        text="Раздел приложения",
        numbered=True,
        number="А.1",
        page=11,
    )
    rendered.post_process(index, document)

    texts = [
        text for paragraph in _toc_paragraphs(rendered) for text in _all_texts(paragraph)
    ]
    assert "ПРИЛОЖЕНИЕ А Исходные данные" in texts
    assert "А.1 " in texts
    assert "Раздел приложения" in texts


def test_content_dot_leader_disabled(ctx, document):
    rendered = render_template("content", {"dot_leader": False}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 0
    _attach_toc_paragraph(rendered, document)
    index = RenderIndex()
    index.add_heading(level=1, text="A", numbered=True, number="1", page=1)
    rendered.post_process(index, document)
    xml = rendered._docx_paragraph._p.xml
    # Tab stops хранят leader как enum — проверяем по w:leader атрибуту.
    assert 'w:leader="dot"' not in xml


def test_content_toc_entries_have_right_dot_leader_tab_stop(ctx, document):
    rendered = render_template("content", {}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 0
    _attach_toc_paragraph(rendered, document)

    index = RenderIndex()
    index.add_heading(level=1, text="A", numbered=True, number="1", page=1)
    index.add_heading(level=2, text="B", numbered=True, number="1.1", page=2)
    rendered.post_process(index, document)

    section = document.sections[0]
    content_width = section.page_width - section.left_margin - section.right_margin
    for paragraph in _toc_paragraphs(rendered):
        tab_stops = _tab_stops(paragraph)
        assert any(
            int(tab.position) == int(content_width)
            and tab.alignment == WD_TAB_ALIGNMENT.RIGHT
            and tab.leader == WD_TAB_LEADER.DOTS
            for tab in tab_stops
        )


def test_content_toc_entries_have_level_indents_and_hanging_indent(ctx, document):
    rendered = render_template("content", {}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 0
    _attach_toc_paragraph(rendered, document)

    index = RenderIndex()
    index.add_heading(
        level=1,
        text="Очень длинное название первого уровня для проверки переносов",
        numbered=True,
        number="1",
        page=1,
    )
    index.add_heading(level=2, text="Второй уровень", numbered=True, number="1.1", page=2)
    index.add_heading(level=3, text="Третий уровень", numbered=True, number="1.1.1", page=3)
    rendered.post_process(index, document)

    paragraphs = _toc_paragraphs(rendered)
    left_indents = [int(p.paragraph_format.left_indent or 0) for p in paragraphs]
    assert left_indents[0] > 0
    assert left_indents[0] < left_indents[1] < left_indents[2]

    for paragraph in paragraphs:
        first_line_indent = paragraph.paragraph_format.first_line_indent
        assert first_line_indent is not None
        assert first_line_indent < 0


def test_content_emits_page_break_after_toc_paragraph(ctx):
    rendered = render_template("content", {}, ctx)
    tracker = LayoutTracker(Length(10_000_000), Length(6_000_000))
    rendered_infos = [
        info for info in rendered.render(None, tracker.current_state)
        if isinstance(info, RenderedInfo)
    ]
    # TOC paragraph + page-break — стартовый параграф из PageBreak.
    last = rendered_infos[-1]
    assert isinstance(last.docx_element.__class__, type)
    assert 'w:type="page"' in last.docx_element._p.xml


def test_content_excludes_headings_on_toc_page(ctx, document):
    """Заголовки на странице TOC (``# *СОДЕРЖАНИЕ``) не должны попадать в оглавление."""

    rendered = render_template("content", {}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 2  # TOC находится на странице 2

    index = RenderIndex()
    index.add_heading(level=1, text="СОДЕРЖАНИЕ", numbered=False, number=None, page=2)
    index.add_heading(level=1, text="Введение", numbered=True, number="1", page=3)
    rendered.post_process(index, document)

    texts = _all_texts(rendered._docx_paragraph)
    assert "СОДЕРЖАНИЕ" not in texts
    assert "Введение" in texts


def test_content_link_is_black_clickable(ctx, document):
    """Ссылка обёрнута в ``<w:hyperlink w:anchor>`` с явным ``w:color=auto``."""

    rendered = render_template("content", {}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 0

    index = RenderIndex()
    entry = index.add_heading(level=1, text="Введение", numbered=True, number="1", page=1)
    rendered.post_process(index, document)

    xml = rendered._docx_paragraph._p.xml
    assert f'w:anchor="{entry.anchor}"' in xml
    assert 'w:val="auto"' in xml


def test_content_emits_default_title_structural_heading(ctx):
    """Шаблон по умолчанию эмитит structural «СОДЕРЖАНИЕ».

    Заголовок выводится отдельным параграфом до основного TOC-параграфа,
    но не проходит через :class:`Numberer` — поэтому ``# Введение`` после
    шаблона остаётся ``"1"``, а не ``"2"``.
    """

    rendered = render_template("content", {}, ctx)
    tracker = LayoutTracker(Length(10_000_000), Length(6_000_000))
    rendered_infos = [
        info for info in rendered.render(None, tracker.current_state)
        if isinstance(info, RenderedInfo)
    ]
    # Первый — title heading; затем TOC paragraph; затем page-break.
    title_p = rendered_infos[0].docx_element
    assert title_p.style.name == "MD2GOST Structural Heading"
    assert "СОДЕРЖАНИЕ" in title_p.text


def test_content_custom_title(ctx):
    rendered = render_template("content", {"title": "Оглавление"}, ctx)
    tracker = LayoutTracker(Length(10_000_000), Length(6_000_000))
    rendered_infos = [
        info for info in rendered.render(None, tracker.current_state)
        if isinstance(info, RenderedInfo)
    ]
    title_p = rendered_infos[0].docx_element
    assert "Оглавление" in title_p.text


def test_content_empty_title_skips_heading(ctx):
    """Пустая строка в ``title`` полностью отключает авто-заголовок."""

    rendered = render_template("content", {"title": ""}, ctx)
    tracker = LayoutTracker(Length(10_000_000), Length(6_000_000))
    rendered_infos = [
        info for info in rendered.render(None, tracker.current_state)
        if isinstance(info, RenderedInfo)
    ]
    # Первый — это уже TOC paragraph (Normal стиль), не Heading 1.
    first = rendered_infos[0].docx_element
    assert first.style.name == "Normal"


def test_content_title_excluded_from_toc_list(ctx, document):
    """Авто-заголовок «СОДЕРЖАНИЕ» не должен попадать в список TOC.

    Заголовок эмитится напрямую через ``DocxParagraph`` и не регистрируется
    в ``RenderIndex``, поэтому в списке его не оказывается естественным
    образом.
    """

    rendered = render_template("content", {}, ctx)
    assert isinstance(rendered, Content)
    rendered._toc_page = 0

    index = RenderIndex()
    index.add_heading(level=1, text="A", numbered=True, number="1", page=1)
    rendered.post_process(index, document)
    texts = _all_texts(rendered._docx_paragraph)
    # ``post_process`` заполняет именно TOC-параграф; заголовок шаблона
    # рендерится отдельным параграфом и в этих ``texts`` не появится.
    assert "СОДЕРЖАНИЕ" not in texts
