from __future__ import annotations

import pytest
from docx import Document
from docx.shared import Cm, Pt

from markdown_gost.render.paragraph_sizer import ParagraphSizer


def _require_font(name: str) -> None:
    """Skip the test if the named font is unavailable on this system."""
    from markdown_gost.render.paragraph_sizer import find_font

    try:
        find_font(name, False, False)
    except (ValueError, OSError, RuntimeError):
        pytest.skip(f"font '{name}' not available in test environment")


def _make_document_with_paragraph(
    text: str,
    *,
    font_name: str = "Times New Roman",
    font_size_pt: int = 14,
    space_before_pt: float = 0,
    space_after_pt: float = 0,
):
    document = Document()
    style = document.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(font_size_pt)

    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(space_before_pt)
    paragraph.paragraph_format.space_after = Pt(space_after_pt)
    return document, paragraph


def test_line_height_for_times_14pt_is_within_tolerance() -> None:
    _require_font("Times New Roman")
    document, paragraph = _make_document_with_paragraph("Привет")
    sizer = ParagraphSizer(paragraph, previous_paragraph=None, max_width=Cm(15))
    result = sizer.calculate_height()

    # Calibration value is 16.05pt; allow ±2pt for cross-platform drift.
    assert result.line_height == pytest.approx(Pt(16.05), abs=Pt(2))


def test_count_lines_short_text() -> None:
    _require_font("Times New Roman")
    document, paragraph = _make_document_with_paragraph("Короткая строка")
    sizer = ParagraphSizer(paragraph, previous_paragraph=None, max_width=Cm(15))
    result = sizer.calculate_height()

    assert result.lines == 1


def test_count_lines_long_text_wraps_multiple_lines() -> None:
    _require_font("Times New Roman")
    long_text = (
        "Это очень длинная строка которая должна перенестись на несколько "
        "строк потому что её ширина больше чем ширина колонки. " * 10
    )
    document, paragraph = _make_document_with_paragraph(long_text)
    # use a narrow column to guarantee wrapping
    sizer = ParagraphSizer(paragraph, previous_paragraph=None, max_width=Cm(8))
    result = sizer.calculate_height()

    assert result.lines > 1


def test_before_after_account_for_spacing() -> None:
    _require_font("Times New Roman")
    document, paragraph = _make_document_with_paragraph(
        "Текст", space_before_pt=12, space_after_pt=6
    )
    sizer = ParagraphSizer(paragraph, previous_paragraph=None, max_width=Cm(15))
    result = sizer.calculate_height()

    assert result.before == Pt(12)
    assert result.after == Pt(6)


def test_find_font_consolas_italic_falls_back_to_oblique_monospace() -> None:
    # Consolas is rarely installed on Linux; the italic fallback target on
    # Debian/Ubuntu is DejaVu Sans Mono Oblique. The fc-list style for that
    # face is "Oblique", not "Italic" — make sure the matcher treats them as
    # interchangeable so the fallback path actually fires.
    from markdown_gost.render.paragraph_sizer import find_font

    try:
        path = find_font("Consolas", bold=False, italic=True)
    except RuntimeError:
        pytest.skip("fc-list (fontconfig) unavailable in test environment")
    except ValueError:
        pytest.skip("no italic/oblique monospace font available in test environment")

    assert path  # the function returned a real path rather than raising


def test_contextual_spacing_uses_previous_after() -> None:
    _require_font("Times New Roman")
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)

    prev = document.add_paragraph("Предыдущий")
    prev.paragraph_format.line_spacing = 1.0
    prev.paragraph_format.space_after = Pt(8)

    curr = document.add_paragraph("Текущий")
    curr.paragraph_format.line_spacing = 1.0
    curr.paragraph_format.space_before = Pt(20)

    # Mark contextual spacing on current paragraph's pPr.
    from docx.oxml import OxmlElement

    ppr = curr._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:contextualSpacing"))

    sizer = ParagraphSizer(curr, previous_paragraph=prev, max_width=Cm(15))
    result = sizer.calculate_height()

    # contextual_spacing + same style => before == previous.after
    assert result.before == Pt(8)
