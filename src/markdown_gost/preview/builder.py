from __future__ import annotations

import logging
import math
import re
from collections.abc import Callable
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Protocol
from urllib.parse import quote

from docx.shared import Length, Pt
from PIL import Image as PILImage
from PIL import UnidentifiedImageError

from markdown_gost.config.schema import CaptionStyle, Config, HeadingLevel, InlineCode
from markdown_gost.config.units import parse_length, parse_pt
from markdown_gost.core.ast import nodes as ast
from markdown_gost.core.parser import parse, walk
from markdown_gost.core.parser.attrs import parse_size_list
from markdown_gost.preview_json import JsonPrimitive, JsonValue, copy_json_layout
from markdown_gost.render.bibliography import BibliographyIndex
from markdown_gost.render.document_factory import _PAGE_SIZES_PORTRAIT
from markdown_gost.render.latex_math import EquationError, latex_to_mathml
from markdown_gost.render.layout_tracker import LayoutTracker
from markdown_gost.render.numberer import Numberer
from markdown_gost.render.references import prepare_document
from markdown_gost.render.render_index import RenderIndex
from markdown_gost.storage.base import Storage
from markdown_gost.storage.fs import FilesystemStorage
from markdown_gost.templates import PreviewTemplateContext, render_preview_template

from .model import (
    PreviewAnchor,
    PreviewBlock,
    PreviewDocument,
    PreviewImageDimension,
    PreviewInline,
    PreviewLength,
    PreviewMargins,
    PreviewPage,
    PreviewPageGeometry,
    PreviewTableCell,
    PreviewTableDimension,
    PreviewTableRow,
)

_EMU_PER_MM = 36_000
_EMU_PER_INCH = 914_400
_CSS_PX_PER_INCH = 96
_MIN_LAYOUT_EMU = 1
_MAX_LIST_LEVEL = 4
_DEFAULT_IMAGE_ASPECT_WIDTH = 16
_DEFAULT_IMAGE_ASPECT_HEIGHT = 9
_HIERARCHICAL_SEPARATOR = "."
_ASSET_ROUTE_PREFIX = "/api/preview/assets"
_DOCX_DEFAULT_PARAGRAPH_SPACE_AFTER = "10pt"
_TABLE_CELL_PAD_LEFT_DXA = 108
_TABLE_CELL_PAD_RIGHT_DXA = 108
_LINE_HEIGHT_CALIBRATION: dict[tuple[str, int], float] = {
    ("Times", 14): 16.13,
    ("Courier", 12): 13.61,
    ("Consolas", 12): 14.75,
    ("Arial", 14): 16.05,
}
_LISTING_TABLE_OVERHEAD_PT = 5.5
_TOC_ENTRY_SPACE_AFTER_PT = 10.0
_CAPTION_LABELS: dict[str, str] = {
    "image": "Рисунок",
    "table": "Таблица",
    "listing": "Листинг",
}
_ALPHABETS_BY_STYLE: dict[str, str] = {
    "lower-alpha-ru": "абвгдежиклмнпрстуфхцшщэюя",
    "upper-alpha-ru": "АБВГДЕЖИКЛМНПРСТУФХЦШЩЭЮЯ",
    "lower-alpha-en": "abcdefghijklmnopqrstuvwxyz",
    "upper-alpha-en": "ABCDEFGHIJKLMNOPQRSTUVWXYZ",
}
_CSS_LENGTH_RE = re.compile(r"^\s*(-?\d+(?:\.\d+)?)([a-zA-Z%]+)\s*$")
_DEFAULT_IMAGE_DPI = 72.0
_LOGGER = logging.getLogger(__name__)


class _ImageMetadataResolver(Protocol):
    def __call__(self, src: str) -> _ImageMetadata | None: ...


@dataclass(frozen=True)
class _PageMetrics:
    width: Length
    height: Length
    top_margin: Length
    right_margin: Length
    bottom_margin: Length
    left_margin: Length
    content_width: Length
    content_height: Length


@dataclass(frozen=True)
class _ImageMetadata:
    width_px: int
    height_px: int
    horizontal_dpi: float = _DEFAULT_IMAGE_DPI
    vertical_dpi: float = _DEFAULT_IMAGE_DPI

    @property
    def native_width(self) -> Length:
        return Pt(self.width_px * 72.0 / self.horizontal_dpi)

    @property
    def native_height(self) -> Length:
        return Pt(self.height_px * 72.0 / self.vertical_dpi)


@dataclass(frozen=True)
class _ChainCtx:
    """Нативная цепочка одного вида — зеркало ``_ChainCtx`` из ``renderable/list.py``.

    Счётчики уровня общие на цепочку (как abstractNum в DOCX): корень даёт
    свежий словарь, продолжение того же вида делит его ссылкой.
    """

    kind: str  # "bullet" | "arabic"
    counters: dict[int, int]
    rel_level: int  # 0-based уровень этого узла цепочки

    def bump(self) -> int:
        """Номер следующего пункта: +1 к своему уровню, ресет глубже лежащих."""
        self.counters[self.rel_level] = self.counters.get(self.rel_level, 0) + 1
        for deeper in [key for key in self.counters if key > self.rel_level]:
            self.counters[deeper] = 0
        return self.counters[self.rel_level]

    def path(self) -> list[int]:
        """Иерархический путь маркера: счётчики уровней 0..rel_level."""
        return [self.counters.get(index, 0) for index in range(self.rel_level + 1)]


def build_preview_model(
    markdown: str,
    config: Config,
    storage: Storage | None = None,
) -> PreviewDocument:
    """Build the internal preview model without DOCX/PDF/LibreOffice.

    ``storage`` is accepted as part of the public preview-build contract. The
    model records asset references only and never embeds binary image data; for
    local filesystem storage it may read image headers to mirror DOCX native
    sizing.
    """

    effective_storage = storage if storage is not None else FilesystemStorage()
    return _PreviewBuilder(config=config, storage=effective_storage).build(markdown)


class _PreviewBuilder:
    def __init__(
        self,
        *,
        config: Config,
        storage: Storage,
    ) -> None:
        self._config = config
        self._storage = storage
        self._metrics = _page_metrics(config)
        self._geometry = _page_geometry(config, self._metrics)
        layout_height = Length(
            max(
                _MIN_LAYOUT_EMU,
                int(self._metrics.content_height) - int(_page_vertical_leading_reserve(config)),
            )
        )
        self._tracker = LayoutTracker(
            layout_height,
            self._metrics.content_width,
        )
        self._numberer = Numberer()
        self._index = RenderIndex()
        self._blocks_by_page: dict[int, list[PreviewBlock]] = {1: []}
        self._paragraph_count = 0
        self._image_count = 0
        self._list_item_count = 0
        self._listing_count = 0
        self._page_break_count = 0
        self._table_count = 0
        self._equation_count = 0
        self._template_block_count = 0
        self._template_deferred: list[tuple[str, Callable[[], None]]] = []
        self._template_blocks: list[tuple[str, PreviewBlock]] = []
        self._template_block_available_heights: dict[int, Length] = {}
        self._image_metadata_cache: dict[str, _ImageMetadata | None] = {}
        self._pending_identifier: str | None = None

    def build(self, markdown: str) -> PreviewDocument:
        document = prepare_document(parse(markdown), self._config)
        bibliography = BibliographyIndex.from_document(document, self._config)
        for item in walk(document):
            inline_children = getattr(item, "children", None)
            if isinstance(inline_children, list):
                item.children = [  # type: ignore[attr-defined]
                    ast.Text(text=bibliography.format_citation(child.key))
                    if isinstance(child, ast.Citation)
                    else child
                    for child in inline_children
                ]
        children = document.children
        index = 0
        while index < len(children):
            node = children[index]
            image = _standalone_image(node.children) if isinstance(node, ast.Paragraph) else None
            self._pending_identifier = image.identifier if image else node.identifier
            if (
                isinstance(node, ast.Caption)
                and node.target == "table"
                and index + 1 < len(children)
                and isinstance(children[index + 1], ast.Table)
            ):
                table = children[index + 1]
                assert isinstance(table, ast.Table)
                self._pending_identifier = table.identifier
                self._add_table(table, caption=node)
                index += 2
                continue
            if (
                isinstance(node, ast.Caption)
                and node.target == "listing"
                and index + 1 < len(children)
                and isinstance(children[index + 1], ast.Listing)
            ):
                listing = children[index + 1]
                assert isinstance(listing, ast.Listing)
                self._pending_identifier = listing.identifier
                self._add_listing(listing, caption=node)
                index += 2
                continue
            if isinstance(node, ast.Heading):
                self._add_heading(node)
            elif isinstance(node, ast.Paragraph):
                self._add_paragraph(node)
            elif isinstance(node, ast.List):
                self._add_list(node)
            elif isinstance(node, ast.Table):
                self._add_table(node, caption=None)
            elif isinstance(node, ast.Listing):
                self._add_listing(node, caption=None)
            elif isinstance(node, ast.Equation):
                self._add_equation(node)
            elif isinstance(node, ast.TemplateBlock):
                self._add_template(node, document)
            elif isinstance(node, ast.PageBreak):
                self._add_page_break()
            elif isinstance(node, ast.AppendixStart):
                letter = self._numberer.start_appendix()
                if self._tracker.current_state.current_page_height > 0:
                    self._tracker.new_page()
                title = f"ПРИЛОЖЕНИЕ {letter}"
                self._add_paragraph(ast.Paragraph(children=[ast.Text(text=title)]))
                if node.title:
                    self._add_paragraph(ast.Paragraph(children=[ast.Text(text=node.title.upper())]))
                self._index.add_heading(
                    level=1,
                    text=f"{title} {node.title or ''}".strip(),
                    numbered=False,
                    number=None,
                    page=self._tracker.current_state.page,
                )
            elif isinstance(node, ast.AppendixEnd):
                self._numberer.end_appendix()
            elif isinstance(node, ast.Bibliography):
                for entry in bibliography.entries():
                    self._add_paragraph(
                        ast.Paragraph(
                            children=[
                                ast.Text(
                                    text=(
                                        f"{entry.number}. "
                                        f"{bibliography.format_entry(entry.source.id)}"
                                    )
                                )
                            ]
                        )
                    )
            elif isinstance(node, ast.Mermaid):
                self._add_paragraph(
                    ast.Paragraph(children=[ast.Text(text="<Mermaid not supported yet>")])
                )
            elif isinstance(node, ast.ThematicBreak):
                page = self._prepare_page_for(Length(12700 * 12))
                self._append_block(
                    PreviewBlock(
                        kind="thematic_break",
                        id=node.identifier or node.node_id,
                        page=page,
                        text="",
                    )
                )
                self._tracker.add_height(Length(12700 * 12))
            index += 1

        for template_name, callback in self._template_deferred:
            try:
                callback()
            except Exception:
                _LOGGER.exception(
                    "Deferred preview rendering failed for template %r",
                    template_name,
                )
                self._add_template_diagnostic(
                    template_name,
                    "deferred preview rendering failed",
                )
        self._paginate_toc_blocks()
        self._finalize_template_layouts()
        total_pages = max(
            1,
            self._tracker.current_state.page,
            *(
                block.page + int(block.kind == "page_break")
                for blocks in self._blocks_by_page.values()
                for block in blocks
            ),
        )
        self._index.set_total_pages(total_pages)
        pages = [
            PreviewPage(number=number, blocks=self._blocks_by_page.get(number, []))
            for number in range(1, total_pages + 1)
        ]
        anchors = [
            PreviewAnchor(
                kind="heading",
                anchor=entry.anchor,
                page=entry.page,
                block_id=entry.anchor,
                text=entry.text,
                level=entry.level,
            )
            for entry in self._index.headings
        ]
        return PreviewDocument(
            geometry=self._geometry,
            total_pages=total_pages,
            pages=pages,
            anchors=anchors,
        )

    def _image_metadata(self, src: str) -> _ImageMetadata | None:
        if src not in self._image_metadata_cache:
            self._image_metadata_cache[src] = _read_filesystem_image_metadata(
                self._storage,
                src,
            )
        return self._image_metadata_cache[src]

    def _add_paragraph(self, node: ast.Paragraph) -> None:
        standalone_image = _standalone_image(node.children)
        if standalone_image is not None:
            self._add_image_block(standalone_image)
            return

        self._paragraph_count += 1
        inlines = _preview_inlines(
            node.children,
            image_metadata=self._image_metadata,
            inline_code=self._config.paragraph.inline_code,
        )
        text = _collect_inline_text(
            node.children,
            inline_code=self._config.paragraph.inline_code,
        )
        height = _estimate_text_height(
            text,
            font_family=self._config.font.family,
            font_size=self._config.font.size,
            line_spacing=self._config.font.line_spacing,
            content_width=self._metrics.content_width,
            space_after=_DOCX_DEFAULT_PARAGRAPH_SPACE_AFTER,
        )
        page = self._prepare_page_for(height)
        block = PreviewBlock(
            kind="paragraph",
            id=f"paragraph-{self._paragraph_count}",
            page=page,
            text=text,
            inlines=inlines,
            style=_paragraph_style(self._config),
        )
        self._append_block(block)
        self._tracker.add_height(height)
        self._ensure_page(self._tracker.current_state.page)

    def _add_image_block(self, node: ast.Image) -> None:
        self._image_count += 1
        number = self._numberer.allocate("image")
        number_text = self._numberer.format_number("image", number)
        caption = _format_caption(
            self._config,
            category="image",
            number=number_text,
            text=node.alt or None,
        )
        image_height = _estimate_image_height(
            node,
            content_width=self._metrics.content_width,
            content_height=self._metrics.content_height,
            metadata=self._image_metadata(node.src),
        )
        caption_height = _estimate_text_height(
            caption,
            font_family=self._config.font.family,
            font_size=self._config.font.size,
            line_spacing=(
                self._config.captions.image.line_spacing or self._config.font.line_spacing
            ),
            content_width=self._metrics.content_width,
            space_before=self._config.captions.image.space_before,
            space_after=self._config.captions.image.space_after,
        )
        height = Length(int(image_height) + int(caption_height))
        page = self._prepare_page_for(height)
        block = PreviewBlock(
            kind="image",
            id=f"image-{self._image_count}",
            page=page,
            text=node.alt,
            inlines=_preview_inline(node, image_metadata=self._image_metadata),
            style=_caption_style(self._config, "image"),
            anchor=f"image-{self._image_count}",
            number=number_text,
            caption=caption,
        )
        self._append_block(block)
        self._tracker.add_height(height)
        self._ensure_page(self._tracker.current_state.page)

    def _add_table(self, node: ast.Table, *, caption: ast.Caption | None) -> None:
        self._table_count += 1
        number = self._numberer.allocate("table")
        number_text = self._numberer.format_number("table", number)
        caption_text = caption.text if caption is not None else None
        rendered_caption = _format_caption(
            self._config,
            category="table",
            number=number_text,
            text=caption_text,
        )
        column_widths, row_heights = _table_dimensions_from_caption(node, caption)
        rows = _preview_table_rows(
            node,
            image_metadata=self._image_metadata,
            inline_code=self._config.paragraph.inline_code,
        )
        text = " ".join(cell.text for row in rows for cell in row.cells if cell.text.strip())
        table_height = _estimate_table_height(
            node,
            config=self._config,
            content_width=self._metrics.content_width,
            row_heights=row_heights,
        )
        caption_height = _estimate_text_height(
            rendered_caption,
            font_family=self._config.font.family,
            font_size=self._config.font.size,
            line_spacing=(
                self._config.captions.table.line_spacing or self._config.font.line_spacing
            ),
            content_width=self._metrics.content_width,
            space_before=self._config.table.space_before,
            space_after=self._config.captions.table.space_after,
        )
        height = Length(int(table_height) + int(caption_height))
        page = self._prepare_page_for(height)
        resolved_column_widths, table_layout = _preview_table_column_geometry(
            node,
            column_widths=column_widths,
            config=self._config,
            content_width=self._metrics.content_width,
        )
        block = PreviewBlock(
            kind="table",
            id=f"table-{self._table_count}",
            page=page,
            text=text,
            style=_table_style(self._config, table_layout=table_layout),
            anchor=f"table-{self._table_count}",
            number=number_text,
            continuation=False,
            caption=rendered_caption,
            rows=rows,
            column_widths=resolved_column_widths,
            row_heights=[_preview_table_dimension(row_height) for row_height in row_heights]
            if row_heights is not None
            else None,
        )
        self._append_block(block)
        self._tracker.add_height(height)
        self._ensure_page(self._tracker.current_state.page)

    def _add_listing(self, node: ast.Listing, *, caption: ast.Caption | None) -> None:
        self._listing_count += 1
        number_text: str | None = None
        rendered_caption: str | None = None
        caption_text = caption.text if caption is not None else None
        caption_height = Length(0)
        if caption is not None:
            number = self._numberer.allocate("listing")
            number_text = self._numberer.format_number("listing", number)
            rendered_caption = _format_caption(
                self._config,
                category="listing",
                number=number_text,
                text=caption_text,
            )
            caption_space_before = self._config.captions.listing.space_before
            if int(parse_length(self._config.listing.space_before)) > 0:
                caption_space_before = self._config.listing.space_before
            caption_height = _estimate_text_height(
                rendered_caption,
                font_family=self._config.font.family,
                font_size=self._config.font.size,
                line_spacing=(
                    self._config.captions.listing.line_spacing or self._config.font.line_spacing
                ),
                content_width=self._metrics.content_width,
                space_before=caption_space_before,
                space_after=self._config.captions.listing.space_after,
            )
        lines = _listing_lines(node.code)
        font_pt = parse_pt(self._config.listing.font.size)
        line_height = Pt(
            _calibrated_line_height_pt(
                self._config.listing.font.family,
                font_pt,
            )
            * self._config.listing.line_spacing
        )
        overhead = Pt(_LISTING_TABLE_OVERHEAD_PT)
        trailing_space = parse_length(self._config.listing.space_after)
        highlighted_extra = (
            line_height if _preview_listing_inlines(node, config=self._config) else Length(0)
        )
        page_font_pt = parse_pt(self._config.font.size)
        footer_clearance = Pt(
            _calibrated_line_height_pt(
                self._config.font.family,
                page_font_pt,
            )
        )
        style = _listing_style(self._config)
        consumed = 0
        chunk_index = 0
        while consumed < len(lines):
            first_chunk = chunk_index == 0
            fixed_height = Length(
                int(overhead) + int(highlighted_extra) + (int(caption_height) if first_chunk else 0)
            )
            required_first_line = Length(int(fixed_height) + int(line_height))
            state = self._tracker.current_state
            usable_remaining = max(
                0,
                int(state.remaining_page_height) - int(footer_clearance),
            )
            if state.current_page_height > 0 and int(required_first_line) > usable_remaining:
                self._tracker.new_page()
                self._ensure_page(self._tracker.current_state.page)
                state = self._tracker.current_state
                usable_remaining = max(
                    0,
                    int(state.remaining_page_height) - int(footer_clearance),
                )

            available_for_lines = max(0, usable_remaining - int(fixed_height))
            capacity = max(1, available_for_lines // max(1, int(line_height)))
            remaining_lines = len(lines) - consumed
            if capacity >= remaining_lines:
                final_available = max(0, available_for_lines - int(trailing_space))
                capacity = max(1, final_available // max(1, int(line_height)))
            chunk_lines = lines[consumed : consumed + capacity]
            consumed += len(chunk_lines)
            has_more = consumed < len(lines)
            chunk_text = _listing_chunk_text(
                chunk_lines,
                trailing_newline=has_more or node.code.endswith("\n"),
            )
            chunk_node = ast.Listing(language=node.language, code=chunk_text)
            page = self._tracker.current_state.page
            block_id = f"listing-{self._listing_count}"
            if not first_chunk:
                block_id = f"{block_id}-continuation-{chunk_index}"
            chunk_style = dict(style)
            chunk_style["space_before"] = "0pt"
            chunk_style["space_after"] = self._config.listing.space_after if not has_more else "0pt"
            if (
                first_chunk
                and rendered_caption
                and int(parse_length(self._config.listing.space_before))
            ):
                chunk_style["caption_space_before"] = self._config.listing.space_before
            block = PreviewBlock(
                kind="listing",
                id=block_id,
                page=page,
                text=chunk_text,
                inlines=_preview_listing_inlines(
                    chunk_node,
                    config=self._config,
                ),
                style=chunk_style,
                anchor=(f"listing-{self._listing_count}" if first_chunk else None),
                number=number_text,
                continuation=not first_chunk,
                caption=rendered_caption if first_chunk else None,
                language=node.language,
            )
            self._append_block(block)
            height = Length(
                int(fixed_height)
                + int(line_height) * len(chunk_lines)
                + (int(trailing_space) if not has_more else 0)
                + (int(footer_clearance) if not has_more else 0)
            )
            self._tracker.add_height(height)
            self._ensure_page(self._tracker.current_state.page)
            chunk_index += 1
            if has_more:
                self._tracker.new_page()
                self._ensure_page(self._tracker.current_state.page)

    def _add_equation(self, node: ast.Equation) -> None:
        self._equation_count += 1
        number = self._numberer.allocate("equation")
        number_text = self._numberer.format_number("equation", number)
        mathml, diagnostic = _equation_mathml(node.latex)
        height = _estimate_equation_height(
            node.latex,
            config=self._config,
            mathml=mathml,
        )
        page = self._prepare_page_for(height)
        block = PreviewBlock(
            kind="equation",
            id=f"equation-{self._equation_count}",
            page=page,
            text=node.latex,
            style=_equation_style(self._config),
            mathml=mathml,
            diagnostic=diagnostic,
            anchor=f"equation-{self._equation_count}",
            number=number_text,
        )
        self._append_block(block)
        self._tracker.add_height(height)
        self._ensure_page(self._tracker.current_state.page)

    def _add_list(
        self,
        node: ast.List,
        *,
        level: int = 1,
        parent_path: list[int] | None = None,
        parent_ctx: _ChainCtx | None = None,
    ) -> None:
        if parent_path is None and level == 1:
            parent_path = []
        clamped_level = min(level, _MAX_LIST_LEVEL)
        counter = node.start if node.ordered else 1
        marker_style = _resolve_list_marker_style(node)
        delimiter = node.delimiter if marker_style == "arabic" else None
        ctx = self._list_chain_ctx(node, marker_style=marker_style, parent_ctx=parent_ctx)
        for item in node.items:
            # Нативная цепочка: номер пункта — счётчик уровня с ресетом глубже.
            counter = ctx.bump() if ctx is not None else counter
            if ctx is not None and ctx.kind == "arabic":
                full_path: list[int] | None = ctx.path()
            elif marker_style == "arabic" and parent_path is not None:
                full_path = [*parent_path, counter]
            else:
                full_path = None
            blocks = _list_item_blocks(item)
            has_para = any(kind == "para" for kind, _ in blocks)
            bullet_rel_level = ctx.rel_level if ctx is not None and ctx.kind == "bullet" else 0
            if not has_para:
                self._add_list_paragraph(
                    [],
                    marker_style=marker_style,
                    counter=counter,
                    full_path=full_path,
                    level=clamped_level,
                    delimiter=delimiter,
                    bullet_rel_level=bullet_rel_level,
                )
            first_para_emitted = False
            for kind, payload in blocks:
                if kind == "para":
                    assert isinstance(payload, list)
                    if not first_para_emitted:
                        first_para_emitted = True
                        self._add_list_paragraph(
                            payload,
                            marker_style=marker_style,
                            counter=counter,
                            full_path=full_path,
                            level=clamped_level,
                            delimiter=delimiter,
                            bullet_rel_level=bullet_rel_level,
                        )
                    else:
                        self._add_list_paragraph(
                            payload,
                            marker_style=marker_style,
                            counter=counter,
                            full_path=full_path,
                            level=clamped_level,
                            delimiter=delimiter,
                            continuation=True,
                        )
                else:
                    assert isinstance(payload, ast.List)
                    self._add_list(
                        payload,
                        level=level + 1,
                        parent_path=full_path,
                        parent_ctx=ctx,
                    )
            counter += 1

    def _list_chain_ctx(
        self,
        node: ast.List,
        *,
        marker_style: str,
        parent_ctx: _ChainCtx | None,
    ) -> _ChainCtx | None:
        """Зеркало ``_chain_ctx`` из ``renderable/list.py``: None → литеральные маркеры."""
        if self._config.lists.mode != "native" or marker_style in _ALPHABETS_BY_STYLE:
            return None
        kind = "bullet" if marker_style == "bullet" else "arabic"
        if parent_ctx is not None and parent_ctx.kind == kind:
            # Продолжение цепочки: те же счётчики, глубже на уровень;
            # start≠1 вложенного узла — startOverride уровня.
            rel_level = parent_ctx.rel_level + 1
            if kind == "arabic" and node.start != 1:
                parent_ctx.counters[rel_level] = node.start - 1
            return _ChainCtx(kind=kind, counters=parent_ctx.counters, rel_level=rel_level)
        counters: dict[int, int] = {}
        if node.ordered and node.start != 1:
            # w:start корня цепочки на ilvl 0.
            counters[0] = node.start - 1
        return _ChainCtx(kind=kind, counters=counters, rel_level=0)

    def _add_list_paragraph(
        self,
        children: list[ast.InlineNode],
        *,
        marker_style: str,
        counter: int,
        full_path: list[int] | None,
        level: int,
        delimiter: str | None,
        bullet_rel_level: int = 0,
        continuation: bool = False,
    ) -> None:
        self._list_item_count += 1
        text = _collect_inline_text(
            children,
            inline_code=self._config.paragraph.inline_code,
        )
        height = _estimate_text_height(
            text,
            font_family=self._config.font.family,
            font_size=self._config.font.size,
            line_spacing=self._config.font.line_spacing,
            content_width=self._metrics.content_width,
        )
        page = self._prepare_page_for(height)
        marker = None
        separator = None
        if not continuation:
            marker = _format_list_marker(
                self._config,
                marker_style=marker_style,
                counter=counter,
                full_path=full_path,
                delimiter=delimiter,
                bullet_rel_level=bullet_rel_level,
            )
            if self._config.lists.mode == "native":
                # Зеркало нумерации DOCX: маркер начинается на left − hanging
                # и занимает ширину hanging, текст — с левого отступа (tab stop
                # = left). Буквальный таб в HTML попадает на стоп 8 пробелов.
                separator = ""
            else:
                separator = "  " if full_path is not None and len(full_path) > 1 else "\t"
        block = PreviewBlock(
            kind="list_item",
            id=f"list-item-{self._list_item_count}",
            page=page,
            text=text,
            inlines=_preview_inlines(
                children,
                image_metadata=self._image_metadata,
                inline_code=self._config.paragraph.inline_code,
            ),
            style=_list_item_style(self._config, level, continuation=continuation),
            level=level,
            list_type="ordered" if marker_style != "bullet" else "unordered",
            marker=marker,
            marker_separator=separator,
            continuation=continuation,
        )
        self._append_block(block)
        self._tracker.add_height(height)
        self._ensure_page(self._tracker.current_state.page)

    def _add_heading(self, node: ast.Heading) -> None:
        raw_text = _collect_inline_text(
            node.children,
            inline_code=self._config.paragraph.inline_code,
        )
        structural = _is_structural_heading(self._config, node, raw_text)
        number = self._next_heading_number(node, structural=structural)
        style = _heading_style(self._config, node, structural=structural)
        height = _estimate_text_height(
            raw_text,
            font_family=self._config.font.family,
            font_size=str(style["font_size"]),
            line_spacing=self._config.font.line_spacing,
            content_width=self._metrics.content_width,
            space_before=str(style["space_before"]),
            space_after=str(style["space_after"]),
        )
        if bool(style["page_break_before"]) and self._tracker.current_state.current_page_height > 0:
            self._tracker.new_page()
            self._ensure_page(self._tracker.current_state.page)

        page = self._prepare_page_for(height)
        entry = self._index.add_heading(
            level=node.level,
            text=raw_text,
            numbered=node.numbered,
            number=number,
            page=page,
        )
        block = PreviewBlock(
            kind="heading",
            id=entry.anchor,
            page=page,
            text=raw_text,
            inlines=_preview_inlines(
                node.children,
                image_metadata=self._image_metadata,
                inline_code=self._config.paragraph.inline_code,
            ),
            style=style,
            anchor=entry.anchor,
            level=node.level,
            numbered=node.numbered,
            number=number,
        )
        self._append_block(block)
        self._tracker.add_height(height)
        self._ensure_page(self._tracker.current_state.page)

    def _add_page_break(self) -> None:
        self._page_break_count += 1
        page = self._tracker.current_state.page
        block = PreviewBlock(
            kind="page_break",
            id=f"page-break-{self._page_break_count}",
            page=page,
            text="",
        )
        self._append_block(block)
        self._tracker.new_page()
        self._ensure_page(self._tracker.current_state.page)

    def _add_template(self, node: ast.TemplateBlock, document: ast.Document) -> None:
        context = PreviewTemplateContext(
            config=self._config,
            document_ast=document,
            index=self._index,
            template_name=node.name,
            _add_block=lambda *args, **kwargs: self._add_template_preview_block(
                node.name,
                *args,
                **kwargs,
            ),
            _add_page_break=self._add_page_break,
            _current_page=lambda: self._tracker.current_state.page,
            _defer=lambda callback: self._template_deferred.append((node.name, callback)),
        )
        diagnostic = render_preview_template(node.name, dict(node.params), context)
        if diagnostic is not None:
            self._add_template_diagnostic(node.name, diagnostic)

    def _add_template_preview_block(
        self,
        template_name: str,
        kind: str,
        text: str,
        *,
        style: dict[str, JsonPrimitive] | None,
        inlines: list[dict[str, Any]] | None,
        block_id: str | None,
        level: int | None,
        numbered: bool | None,
        number: str | None,
        anchor: str | None,
        layout: dict[str, JsonValue] | None,
    ) -> None:
        self._template_block_count += 1
        effective_style = style or _paragraph_style(self._config)
        materialized_inlines = _materialize_template_inlines(inlines)
        height = _estimate_text_height(
            text,
            font_family=str(effective_style.get("font_family", self._config.font.family)),
            font_size=str(effective_style.get("font_size", self._config.font.size)),
            line_spacing=_template_line_spacing(
                effective_style,
                self._config.font.line_spacing,
            ),
            content_width=self._metrics.content_width,
            space_before=str(effective_style.get("space_before", "0pt")),
            space_after=str(effective_style.get("space_after", "0pt")),
        )
        if (
            bool(effective_style.get("page_break_before"))
            and self._tracker.current_state.current_page_height > 0
        ):
            self._tracker.new_page()
            self._ensure_page(self._tracker.current_state.page)
        page = self._prepare_page_for(height)
        block = PreviewBlock(
            kind=kind,
            id=block_id or f"template-block-{self._template_block_count}",
            page=page,
            text=text,
            inlines=materialized_inlines,
            style=effective_style,
            anchor=anchor,
            level=level,
            numbered=numbered,
            number=number,
            layout=layout,
        )
        self._append_block(block)
        self._template_blocks.append((template_name, block))
        if kind == "toc":
            self._template_block_available_heights[id(block)] = (
                self._tracker.current_state.remaining_page_height
            )
        self._tracker.add_height(height)
        self._ensure_page(self._tracker.current_state.page)

    def _add_template_diagnostic(self, name: str, reason: str) -> None:
        self._template_block_count += 1
        text = f"[template '{name}' error: {reason}]"
        height = _estimate_text_height(
            text,
            font_family=self._config.font.family,
            font_size=self._config.font.size,
            line_spacing=self._config.font.line_spacing,
            content_width=self._metrics.content_width,
        )
        page = self._prepare_page_for(height)
        self._append_block(
            PreviewBlock(
                kind="template_diagnostic",
                id=f"template-diagnostic-{self._template_block_count}",
                page=page,
                text=text,
                style=_paragraph_style(self._config),
            )
        )
        self._tracker.add_height(height)
        self._ensure_page(self._tracker.current_state.page)

    def _finalize_template_layouts(self) -> None:
        """Revalidate layouts after template deferred callbacks have run."""

        for template_name, block in self._template_blocks:
            try:
                layout = copy_json_layout(block.layout)
            except Exception:
                _LOGGER.exception(
                    "Final preview layout validation failed for template %r",
                    template_name,
                )
                self._replace_template_block(
                    block,
                    self._template_layout_diagnostic(block.page),
                )
                continue
            self._replace_template_block(block, replace(block, layout=layout))

    def _paginate_toc_blocks(self) -> None:
        """Split finalized semantic TOCs and shift subsequent preview pages.

        Template callbacks populate TOC entries only after the initial layout
        has assigned heading pages.  This narrow second pass reserves those
        rows, inserts continuation pages as needed, then uses block anchors as
        the single source of truth to refresh the render index and TOC links.
        """

        template_index = 0
        while template_index < len(self._template_blocks):
            template_name, block = self._template_blocks[template_index]
            if block.kind != "toc" or block.layout is None:
                template_index += 1
                continue
            try:
                layout = copy_json_layout(block.layout)
            except Exception:
                template_index += 1
                continue
            if layout is None:  # pragma: no cover - guarded by block.layout above.
                template_index += 1
                continue
            entries = layout.get("entries")
            if not isinstance(entries, list):
                template_index += 1
                continue
            chunks = self._toc_entry_chunks(
                block,
                entries,
                first_page_height=self._template_block_available_heights.get(id(block)),
            )
            if len(chunks) <= 1:
                template_index += 1
                continue
            continuations = [
                replace(
                    block,
                    id=(block.id if chunk_index == 0 else f"{block.id}-continuation-{chunk_index}"),
                    page=block.page + chunk_index,
                    continuation=chunk_index > 0,
                    layout={**layout, "entries": chunk},
                )
                for chunk_index, chunk in enumerate(chunks)
            ]
            self._insert_toc_continuations(
                template_index,
                block,
                continuations,
                template_name,
            )
            template_index += len(continuations)

        self._refresh_index_and_toc_pages()

    def _toc_entry_chunks(
        self,
        block: PreviewBlock,
        entries: list[JsonValue],
        *,
        first_page_height: Length | None,
    ) -> list[list[JsonValue]]:
        if not entries:
            return [entries]
        full_page_height = self._tracker.current_state.max_height
        remaining = full_page_height if first_page_height is None else first_page_height
        chunks: list[list[JsonValue]] = [[]]
        for entry in entries:
            row_height = self._toc_entry_height(block, entry)
            if row_height > remaining and (chunks[-1] or remaining < full_page_height):
                chunks.append([])
                remaining = full_page_height
            chunks[-1].append(entry)
            remaining = Length(max(0, int(remaining) - int(row_height)))
        return chunks

    def _toc_entry_height(self, block: PreviewBlock, entry: JsonValue) -> Length:
        if not isinstance(entry, dict):
            return Length(
                int(_estimate_text_height(
                    " ",
                    font_family=self._config.font.family,
                    font_size=self._config.font.size,
                    line_spacing=self._config.font.line_spacing,
                    content_width=self._metrics.content_width,
                ))
                + int(Pt(_TOC_ENTRY_SPACE_AFTER_PT))
            )
        level = entry.get("level", 1)
        safe_level = level if isinstance(level, int) and not isinstance(level, bool) else 1
        indentation = parse_length("0.75cm") * max(0, safe_level - 1)
        content_width = Length(max(_MIN_LAYOUT_EMU, int(self._metrics.content_width) - indentation))
        text = " ".join(str(entry.get(field) or "") for field in ("number", "text", "page"))
        line_height = _estimate_text_height(
            text,
            font_family=str(block.style.get("font_family", self._config.font.family)),
            font_size=str(block.style.get("font_size", self._config.font.size)),
            line_spacing=_template_line_spacing(block.style, self._config.font.line_spacing),
            content_width=content_width,
        )
        # Зеркало CSS превью: .md2gost-toc-entry { margin-bottom: 10pt }.
        return Length(int(line_height) + int(Pt(_TOC_ENTRY_SPACE_AFTER_PT)))

    def _insert_toc_continuations(
        self,
        template_index: int,
        original: PreviewBlock,
        continuations: list[PreviewBlock],
        template_name: str,
    ) -> None:
        """Replace one TOC block while preserving the order of later blocks."""

        ordered_blocks = [
            block for page in sorted(self._blocks_by_page) for block in self._blocks_by_page[page]
        ]
        original_index = next(
            index for index, block in enumerate(ordered_blocks) if block is original
        )
        extra_pages = len(continuations) - 1
        replacements: dict[int, PreviewBlock] = {}
        rebuilt_blocks = [*ordered_blocks[:original_index], *continuations]
        for block in ordered_blocks[original_index + 1 :]:
            shifted = replace(block, page=block.page + extra_pages)
            replacements[id(block)] = shifted
            available_height = self._template_block_available_heights.pop(id(block), None)
            if available_height is not None:
                self._template_block_available_heights[id(shifted)] = available_height
            rebuilt_blocks.append(shifted)

        self._index.numbered_objects = {
            category: [
                replace(entry, page=entry.page + extra_pages)
                if entry.page > original.page
                else entry
                for entry in entries
            ]
            for category, entries in self._index.numbered_objects.items()
        }

        self._blocks_by_page = {}
        for block in rebuilt_blocks:
            self._append_block(block)

        rebuilt_templates: list[tuple[str, PreviewBlock]] = []
        for index, (name, block) in enumerate(self._template_blocks):
            if index == template_index:
                rebuilt_templates.extend((template_name, block) for block in continuations)
            else:
                rebuilt_templates.append((name, replacements.get(id(block), block)))
        self._template_blocks = rebuilt_templates

    def _refresh_index_and_toc_pages(self) -> None:
        """Synchronize render-index pages and finalized TOC links with blocks."""

        pages_by_anchor = {
            block.anchor: block.page
            for blocks in self._blocks_by_page.values()
            for block in blocks
            if block.anchor is not None
        }
        self._index.headings = [
            replace(entry, page=pages_by_anchor.get(entry.anchor, entry.page))
            for entry in self._index.headings
        ]
        self._index.numbered_objects = {
            category: [
                replace(entry, page=pages_by_anchor.get(entry.anchor, entry.page))
                for entry in entries
            ]
            for category, entries in self._index.numbered_objects.items()
        }
        for _, block in self._template_blocks:
            if block.kind != "toc" or block.layout is None:
                continue
            entries = block.layout.get("entries")
            if not isinstance(entries, list):
                continue
            for entry in entries:
                if not isinstance(entry, dict):
                    continue
                anchor = entry.get("anchor")
                if isinstance(anchor, str) and anchor in pages_by_anchor:
                    entry["page"] = pages_by_anchor[anchor]

    def _replace_template_block(self, original: PreviewBlock, replacement: PreviewBlock) -> None:
        for blocks in self._blocks_by_page.values():
            for index, block in enumerate(blocks):
                if block is original:
                    blocks[index] = replacement
                    return
        raise RuntimeError("template preview block missing from page")  # pragma: no cover

    def _template_layout_diagnostic(self, page: int) -> PreviewBlock:
        self._template_block_count += 1
        return PreviewBlock(
            kind="template_diagnostic",
            id=f"template-diagnostic-{self._template_block_count}",
            page=page,
            text="[template error: preview rendering failed]",
            style=_paragraph_style(self._config),
        )

    def _prepare_page_for(self, height: Length) -> int:
        if (
            self._tracker.current_state.current_page_height > 0
            and not self._tracker.can_fit_to_page(height)
        ):
            self._tracker.new_page()
            self._ensure_page(self._tracker.current_state.page)
        return self._tracker.current_state.page

    def _append_block(self, block: PreviewBlock) -> None:
        if self._pending_identifier:
            block = replace(block, id=self._pending_identifier)
            self._pending_identifier = None
        self._ensure_page(block.page)
        self._blocks_by_page[block.page].append(block)

    def _ensure_page(self, page: int) -> None:
        for number in range(1, page + 1):
            self._blocks_by_page.setdefault(number, [])

    def _next_heading_number(
        self,
        node: ast.Heading,
        *,
        structural: bool,
    ) -> str | None:
        if structural:
            return None
        if self._config.headings.numbering == "none":
            return None
        prefix = self._numberer.bump_heading(node.level)
        if not node.numbered:
            return None
        return prefix


def _page_metrics(config: Config) -> _PageMetrics:
    width, height = _PAGE_SIZES_PORTRAIT[config.page.size]
    if config.page.orientation == "landscape":
        width, height = height, width
    top = parse_length(config.page.margins.top)
    right = parse_length(config.page.margins.right)
    bottom = parse_length(config.page.margins.bottom)
    left = parse_length(config.page.margins.left)
    content_width = Length(max(_MIN_LAYOUT_EMU, int(width) - int(left) - int(right)))
    content_height = Length(max(_MIN_LAYOUT_EMU, int(height) - int(top) - int(bottom)))
    return _PageMetrics(
        width=width,
        height=height,
        top_margin=top,
        right_margin=right,
        bottom_margin=bottom,
        left_margin=left,
        content_width=content_width,
        content_height=content_height,
    )


def _page_geometry(config: Config, metrics: _PageMetrics) -> PreviewPageGeometry:
    return PreviewPageGeometry(
        page_size=config.page.size,
        orientation=config.page.orientation,
        width=_preview_length(metrics.width),
        height=_preview_length(metrics.height),
        margins=PreviewMargins(
            top=_preview_length(metrics.top_margin),
            right=_preview_length(metrics.right_margin),
            bottom=_preview_length(metrics.bottom_margin),
            left=_preview_length(metrics.left_margin),
        ),
        content_width=_preview_length(metrics.content_width),
        content_height=_preview_length(metrics.content_height),
    )


def _preview_length(length: Length) -> PreviewLength:
    value = round(int(length) / _EMU_PER_MM, 4)
    return PreviewLength(value=value, unit="mm", emu=int(length))


def _page_vertical_leading_reserve(config: Config) -> Length:
    """Reserve the two half-leading regions LibreOffice keeps at page edges."""

    font_pt = parse_pt(config.font.size)
    line_height_pt = _calibrated_line_height_pt(config.font.family, font_pt)
    leading_pt = max(0.0, line_height_pt * config.font.line_spacing - font_pt)
    return Pt(leading_pt)


def _estimate_text_height(
    text: str,
    *,
    font_family: str,
    font_size: str,
    line_spacing: float,
    content_width: Length,
    space_before: str = "0pt",
    space_after: str = "0pt",
) -> Length:
    font_pt = parse_pt(font_size)
    width_mm = max(1.0, int(content_width) / _EMU_PER_MM)
    avg_char_width_mm = max(0.1, font_pt * 0.1764)
    chars_per_line = max(1, int(width_mm / avg_char_width_mm))
    visual_text = text or " "
    lines = 0
    for part in visual_text.splitlines() or [""]:
        lines += max(1, math.ceil(len(part) / chars_per_line))
    line_height = Pt(_calibrated_line_height_pt(font_family, font_pt) * line_spacing)
    before = parse_length(space_before)
    after = parse_length(space_after)
    return Length(max(_MIN_LAYOUT_EMU, int(line_height) * lines + int(before) + int(after)))


def _template_line_spacing(style: dict[str, Any], default: float) -> float:
    value: object = style.get("line_spacing", default)
    if not isinstance(value, str | int | float):
        return default
    try:
        return float(value)
    except ValueError:
        return default


def _calibrated_line_height_pt(font_family: str, font_size_pt: float) -> float:
    rounded = int(round(font_size_pt))
    for (name_fragment, size), line_height in _LINE_HEIGHT_CALIBRATION.items():
        if size == rounded and name_fragment in font_family:
            return line_height
    return font_size_pt


def _paragraph_style(config: Config) -> dict[str, JsonPrimitive]:
    return {
        "font_family": config.font.family,
        "font_size": config.font.size,
        "line_spacing": config.font.line_spacing,
        "alignment": config.paragraph.alignment,
        "indent_first_line": config.paragraph.indent_first_line,
        "space_after": _DOCX_DEFAULT_PARAGRAPH_SPACE_AFTER,
    }


def _list_item_style(
    config: Config,
    level: int,
    *,
    continuation: bool,
) -> dict[str, JsonPrimitive]:
    text_indent = "0cm" if continuation else _negative_css_length(config.lists.indent_per_level)
    style: dict[str, JsonPrimitive] = {
        "font_family": config.font.family,
        "font_size": config.font.size,
        "line_spacing": config.font.line_spacing,
        "alignment": config.paragraph.alignment,
        "margin_left": _list_margin_left(config, level),
        "text_indent": text_indent,
        "space_before": "0pt",
        "space_after": "0pt",
    }
    if config.lists.mode == "native" and not continuation:
        # Ширина блока маркера = hanging (смещение первой строки), как в
        # нумерации DOCX: маркер на left − hanging, текст на left.
        style["marker_width"] = config.lists.indent_per_level
    return style


def _table_style(config: Config, *, table_layout: str) -> dict[str, JsonPrimitive]:
    style: dict[str, JsonPrimitive] = {
        "font_family": config.font.family,
        "font_size": config.table.font_size or config.font.size,
        "line_spacing": config.font.line_spacing,
        "space_before": config.table.space_before,
        "space_after": config.table.space_after,
        "caption_alignment": config.captions.table.alignment,
        "table_layout": table_layout,
        "repeat_header_on_break": config.table.repeat_header_on_break,
        "header_bold": config.table.header_bold,
    }
    style.update(_prefixed_caption_style(config, "table"))
    return style


def _listing_style(config: Config) -> dict[str, JsonPrimitive]:
    style: dict[str, JsonPrimitive] = {
        "font_family": config.listing.font.family,
        "font_size": config.listing.font.size,
        "line_spacing": config.listing.line_spacing,
        "space_before": config.listing.space_before,
        "space_after": config.listing.space_after,
        "caption_alignment": config.captions.listing.alignment,
    }
    style.update(_prefixed_caption_style(config, "listing"))
    return style


def _equation_style(config: Config) -> dict[str, JsonPrimitive]:
    return {
        "font_family": config.font.family,
        "font_size": config.font.size,
        "line_spacing": config.font.line_spacing,
        "alignment": "center",
        # DOCX uses exact-height spacer rows inside the equation table.
        # Padding preserves that additive geometry; adjacent CSS margins would
        # collapse with paragraph spacing and progressively compress the page.
        "padding_before": config.equation.space_before,
        "padding_after": config.equation.space_after,
        "space_before": "0pt",
        "space_after": "0pt",
        "numbering_alignment": config.equation.numbering_alignment,
        "parentheses": config.equation.parentheses,
    }


def _list_margin_left(config: Config, level: int) -> str:
    base = _parse_css_length(config.lists.indent_left)
    per_level = _parse_css_length(config.lists.indent_per_level)
    if base is not None and per_level is not None and base[1] == per_level[1]:
        value = base[0] + per_level[0] * (level - 1)
        return _format_css_length(value, base[1])
    left = parse_length(config.lists.indent_left)
    step = parse_length(config.lists.indent_per_level)
    return _length_to_mm_css(Length(int(left) + int(step) * (level - 1)))


def _negative_css_length(value: str) -> str:
    parsed = _parse_css_length(value)
    if parsed is None:
        return f"-{value}"
    return _format_css_length(-parsed[0], parsed[1])


def _parse_css_length(value: str) -> tuple[float, str] | None:
    match = _CSS_LENGTH_RE.match(value)
    if match is None:
        return None
    return float(match.group(1)), match.group(2)


def _format_css_length(value: float, unit: str) -> str:
    formatted = f"{value:.4f}".rstrip("0").rstrip(".")
    if formatted == "-0":
        formatted = "0"
    return f"{formatted}{unit}"


def _length_to_mm_css(length: Length) -> str:
    return _format_css_length(int(length) / _EMU_PER_MM, "mm")


def _heading_style(
    config: Config,
    node: ast.Heading,
    *,
    structural: bool,
) -> dict[str, JsonPrimitive]:
    spec = (
        config.headings.structural
        if structural
        else config.headings.levels.get(node.level, HeadingLevel())
    )
    return {
        "font_family": config.font.family,
        "font_size": spec.size or config.font.size,
        "line_spacing": config.font.line_spacing,
        "alignment": spec.alignment,
        "indent_first_line": "0cm",
        "space_before": spec.space_before,
        "space_after": spec.space_after,
        "page_break_before": spec.page_break_before,
        "keep_with_next": spec.keep_with_next,
        "bold": spec.bold,
        "italic": spec.italic,
        "uppercase": spec.uppercase,
    }


def _caption_style(config: Config, category: str) -> dict[str, JsonPrimitive]:
    spec = _caption_spec(config, category)
    return {
        "font_family": config.font.family,
        "font_size": config.font.size,
        "line_spacing": spec.line_spacing or config.font.line_spacing,
        "alignment": spec.alignment,
        "space_before": spec.space_before,
        "space_after": spec.space_after,
        "bold": spec.bold,
        "italic": spec.italic,
    }


def _prefixed_caption_style(
    config: Config,
    category: str,
) -> dict[str, JsonPrimitive]:
    return {f"caption_{name}": value for name, value in _caption_style(config, category).items()}


_TEMPLATE_INLINE_STRING_FIELDS = frozenset(
    {
        "text",
        "color",
        "font_family",
        "font_size",
        "mathml",
        "diagnostic",
        "href",
        "src",
        "asset_path",
        "alt",
        "title",
        "aspect_ratio",
    }
)
_TEMPLATE_INLINE_BOOLEAN_FIELDS = frozenset({"bold", "italic", "strike", "underline"})
_TEMPLATE_INLINE_DIMENSION_FIELDS = frozenset({"width", "height"})
_TEMPLATE_INLINE_FIELDS = (
    {"kind"}
    | _TEMPLATE_INLINE_STRING_FIELDS
    | _TEMPLATE_INLINE_BOOLEAN_FIELDS
    | _TEMPLATE_INLINE_DIMENSION_FIELDS
)


def _materialize_template_inlines(
    inlines: list[dict[str, Any]] | None,
) -> list[PreviewInline]:
    """Validate and detach extension inlines before adding them to the model."""

    if inlines is None:
        return []
    if not isinstance(inlines, list):
        raise ValueError("template inlines must be a list")
    return [_materialize_template_inline(inline) for inline in inlines]


def _materialize_template_inline(inline: object) -> PreviewInline:
    if not isinstance(inline, dict):
        raise ValueError("template inline must be an object")
    if not all(isinstance(key, str) for key in inline):
        raise ValueError("template inline keys must be strings")
    unknown_fields = set(inline) - _TEMPLATE_INLINE_FIELDS
    if unknown_fields:
        raise ValueError("template inline contains unsupported fields")

    kind = inline.get("kind")
    if not isinstance(kind, str):
        raise ValueError("template inline kind must be a string")
    values: dict[str, Any] = {"kind": kind}
    for field_name in _TEMPLATE_INLINE_STRING_FIELDS:
        value = inline.get(field_name)
        if value is not None:
            if not isinstance(value, str):
                raise ValueError(f"template inline {field_name} must be a string")
            values[field_name] = value
    for field_name in _TEMPLATE_INLINE_BOOLEAN_FIELDS:
        value = inline.get(field_name)
        if value is not None:
            if not isinstance(value, bool):
                raise ValueError(f"template inline {field_name} must be a boolean")
            values[field_name] = value
    for field_name in _TEMPLATE_INLINE_DIMENSION_FIELDS:
        value = inline.get(field_name)
        if value is not None:
            values[field_name] = _materialize_template_image_dimension(value)
    return PreviewInline(**values)


def _materialize_template_image_dimension(value: object) -> PreviewImageDimension:
    if not isinstance(value, dict) or set(value) != {"value", "unit"}:
        raise ValueError("template image dimension must have value and unit")
    raw_value = value["value"]
    unit = value["unit"]
    if isinstance(raw_value, bool) or not isinstance(raw_value, int | float):
        raise ValueError("template image dimension value must be numeric")
    if not isinstance(unit, str):
        raise ValueError("template image dimension unit must be a string")
    numeric_value = float(raw_value)
    if not math.isfinite(numeric_value):
        raise ValueError("template image dimension value must be finite")
    return PreviewImageDimension(value=numeric_value, unit=unit)


def _preview_inlines(
    children: list[ast.InlineNode],
    *,
    image_metadata: _ImageMetadataResolver | None = None,
    inline_code: InlineCode | None = None,
    bold: bool = False,
    italic: bool = False,
    strike: bool = False,
    underline: bool = False,
) -> list[PreviewInline]:
    out: list[PreviewInline] = []
    for child in children:
        out.extend(
            _preview_inline(
                child,
                image_metadata=image_metadata,
                inline_code=inline_code,
                bold=bold,
                italic=italic,
                strike=strike,
                underline=underline,
            )
        )
    return out


def _preview_inline(
    node: ast.InlineNode,
    *,
    image_metadata: _ImageMetadataResolver | None = None,
    inline_code: InlineCode | None = None,
    bold: bool = False,
    italic: bool = False,
    strike: bool = False,
    underline: bool = False,
) -> list[PreviewInline]:
    if isinstance(node, ast.Text):
        return [
            PreviewInline(
                kind="text",
                text=node.text,
                bold=bold,
                italic=italic,
                strike=strike,
                underline=underline,
            )
        ]
    if isinstance(node, ast.InlineCode):
        spec = inline_code or InlineCode()
        text = node.code
        if spec.quotes:
            text = f"«{text}»"
        text = text.replace("-", "‑")
        return [
            PreviewInline(
                kind="code",
                text=text,
                bold=bold,
                italic=italic or spec.italic,
                strike=strike,
                underline=underline,
                font_family=str(spec.font),
                font_size=spec.size,
            )
        ]
    if isinstance(node, ast.LineBreak):
        return [PreviewInline(kind="line_break", text=" ")]
    if isinstance(node, ast.Image):
        metadata = image_metadata(node.src) if image_metadata is not None else None
        infer_native_size = node.width is None and node.height is None
        native_metadata = metadata if infer_native_size else None
        return [
            PreviewInline(
                kind="image",
                src=_asset_url(node.src),
                asset_path=node.src,
                alt=node.alt,
                title=node.title,
                width=_preview_image_width_dimension(
                    node.width,
                    metadata=native_metadata,
                ),
                height=_preview_image_height_dimension(
                    node.height,
                    metadata=native_metadata,
                ),
                aspect_ratio=_image_aspect_ratio(node, metadata=metadata),
            )
        ]
    if isinstance(node, ast.Link):
        text = _collect_inline_text(node.children, inline_code=inline_code)
        return [
            PreviewInline(
                kind="link",
                text=text,
                href=node.url,
                bold=bold,
                italic=italic,
                strike=strike,
                underline=underline,
            )
        ]
    if isinstance(node, ast.Reference):
        raise ValueError(f"unresolved reference: {node.name}")
    if isinstance(node, ast.Citation):
        raise ValueError(f"unresolved citation: {node.key}")
    if isinstance(
        node,
        ast.Emphasis | ast.Strong | ast.Strikethrough | ast.Underline,
    ):
        return _preview_inlines(
            node.children,
            image_metadata=image_metadata,
            inline_code=inline_code,
            bold=bold or isinstance(node, ast.Strong),
            italic=italic or isinstance(node, ast.Emphasis),
            strike=strike or isinstance(node, ast.Strikethrough),
            underline=underline or isinstance(node, ast.Underline),
        )
    if isinstance(node, ast.InlineEquation):
        mathml, diagnostic = _equation_mathml(node.latex)
        return [
            PreviewInline(
                kind="inline_equation",
                text=node.latex,
                bold=bold,
                italic=italic,
                strike=strike,
                underline=underline,
                mathml=mathml,
                diagnostic=diagnostic,
            )
        ]
    return [
        PreviewInline(
            kind="text",
            text=_collect_inline_text([node], inline_code=inline_code),
            bold=bold,
            italic=italic,
            strike=strike,
            underline=underline,
        )
    ]


def _preview_image_width_dimension(
    spec: ast.Length | None,
    *,
    metadata: _ImageMetadata | None = None,
) -> PreviewImageDimension | None:
    if spec is None:
        if metadata is None:
            return None
        return PreviewImageDimension(
            value=_length_to_points(metadata.native_width),
            unit="pt",
        )
    if spec.unit == "auto":
        return None
    return PreviewImageDimension(value=float(spec.value), unit=spec.unit)


def _preview_image_height_dimension(
    spec: ast.Length | None,
    *,
    metadata: _ImageMetadata | None = None,
) -> PreviewImageDimension | None:
    if spec is None:
        if metadata is None:
            return None
        return PreviewImageDimension(
            value=_length_to_points(metadata.native_height),
            unit="pt",
        )
    if spec.unit == "auto":
        return None
    return PreviewImageDimension(value=float(spec.value), unit=spec.unit)


def _preview_table_dimension(spec: ast.Length) -> PreviewTableDimension:
    return PreviewTableDimension(value=float(spec.value), unit=spec.unit)


def _preview_table_rows(
    node: ast.Table,
    *,
    image_metadata: _ImageMetadataResolver | None = None,
    inline_code: InlineCode | None = None,
) -> list[PreviewTableRow]:
    rows: list[PreviewTableRow] = []
    for row in node.rows:
        cells = [
            PreviewTableCell(
                text=_collect_inline_text(cell.children, inline_code=inline_code),
                inlines=_preview_inlines(
                    cell.children,
                    image_metadata=image_metadata,
                    inline_code=inline_code,
                ),
                align=cell.align,
                header=cell.header,
            )
            for cell in row.cells
        ]
        rows.append(PreviewTableRow(cells=cells, header=row.header))
    return rows


def _preview_listing_inlines(
    node: ast.Listing,
    *,
    config: Config,
) -> list[PreviewInline]:
    if not config.listing.syntax_highlighting or not node.language:
        return []
    text = (node.code or "").rstrip("\n")
    try:
        from pygments import lex
        from pygments.lexers import get_lexer_by_name
        from pygments.styles import get_style_by_name
        from pygments.util import ClassNotFound
    except ImportError:  # pragma: no cover - pygments is a runtime dependency
        return []

    try:
        lexer = get_lexer_by_name(node.language)
        style = get_style_by_name("sas")
    except ClassNotFound:
        return []

    token_styles = {token: style.style_for_token(token) for token, _ in style}
    out: list[PreviewInline] = []
    for token_type, value in lex(text, lexer):
        token_style = token_styles.get(token_type) or {}
        current = token_type
        while not token_style and current.parent is not None:
            current = current.parent
            token_style = token_styles.get(current) or token_style
        color = token_style.get("color")
        out.append(
            PreviewInline(
                kind="listing_token",
                text=value,
                bold=bool(token_style.get("bold")),
                italic=bool(token_style.get("italic")),
                color=f"#{color}" if color else None,
            )
        )
    return out


def _preview_table_column_geometry(
    node: ast.Table,
    *,
    column_widths: list[ast.Length] | None,
    config: Config,
    content_width: Length,
) -> tuple[list[PreviewTableDimension] | None, str]:
    ncols = max((len(row.cells) for row in node.rows), default=0)
    if ncols == 0:
        return None, "autofit"
    if column_widths is not None and any(width.unit != "auto" for width in column_widths):
        dxas = _resolve_fixed_table_column_dxas(
            column_widths,
            content_width=content_width,
        )
        return [_preview_dxa_dimension(dxa) for dxa in dxas], "fixed"
    dxas = _estimate_autofit_table_column_dxas(
        node,
        ncols=ncols,
        config=config,
        content_width=content_width,
    )
    return [_preview_dxa_dimension(dxa) for dxa in dxas], "autofit"


def _resolve_fixed_table_column_dxas(
    column_widths: list[ast.Length],
    *,
    content_width: Length,
) -> list[int]:
    content_dxa = _emu_to_dxa(content_width)
    explicit: list[int | None] = []
    explicit_total = 0
    auto_count = 0
    for spec in column_widths:
        dxa = _table_length_to_dxa(spec, percent_base=content_width)
        explicit.append(dxa)
        if dxa is None:
            auto_count += 1
        else:
            explicit_total += dxa
    if auto_count == 0:
        return [int(dxa) for dxa in explicit if dxa is not None]
    remainder = content_dxa - explicit_total
    if remainder <= 0:
        raise ValueError(
            "widths: explicit columns exceed table content width; reduce explicit values"
        )
    per_auto = remainder // auto_count
    return [dxa if dxa is not None else per_auto for dxa in explicit]


def _estimate_autofit_table_column_dxas(
    node: ast.Table,
    *,
    ncols: int,
    config: Config,
    content_width: Length,
) -> list[int]:
    font_pt = parse_pt(config.table.font_size or config.font.size)
    char_dxa = max(int(font_pt * 11), 100)
    cell_padding = _TABLE_CELL_PAD_LEFT_DXA + _TABLE_CELL_PAD_RIGHT_DXA
    min_col = char_dxa * 2 + cell_padding

    natural: list[int] = [min_col] * ncols
    for row in node.rows:
        for col_idx, cell in enumerate(row.cells[:ncols]):
            text_len = len(
                _collect_inline_text(
                    cell.children,
                    inline_code=config.paragraph.inline_code,
                )
            )
            natural[col_idx] = max(natural[col_idx], text_len * char_dxa + cell_padding)

    content_dxa = _emu_to_dxa(content_width)
    total = sum(natural)
    if total <= 0:
        per_column = content_dxa // ncols
        return [per_column] * ncols
    scale = content_dxa / total
    widths = [max(int(round(width * scale)), min_col) for width in natural]
    diff = content_dxa - sum(widths)
    if diff:
        widest_idx = max(range(ncols), key=lambda idx: widths[idx])
        widths[widest_idx] += diff
    return widths


def _table_length_to_dxa(spec: ast.Length, *, percent_base: Length) -> int | None:
    if spec.unit == "auto":
        return None
    if spec.unit == "%":
        return _emu_to_dxa(Length(int(int(percent_base) * spec.value / 100.0)))
    return _emu_to_dxa(parse_length(f"{spec.value}{spec.unit}"))


def _emu_to_dxa(length: Length) -> int:
    return int(round(int(length) / _EMU_PER_INCH * 1440))


def _preview_dxa_dimension(dxa: int) -> PreviewTableDimension:
    return PreviewTableDimension(value=round(dxa / 20.0, 2), unit="pt")


def _table_dimensions_from_caption(
    node: ast.Table,
    caption: ast.Caption | None,
) -> tuple[list[ast.Length] | None, list[ast.Length] | None]:
    column_widths = node.column_widths
    row_heights = node.row_heights
    if caption is None:
        return column_widths, row_heights

    ncols = max((len(row.cells) for row in node.rows), default=0)
    nrows = len(node.rows)
    raw_widths = caption.attrs.get("widths")
    if raw_widths:
        pairs = parse_size_list(
            raw_widths,
            expected_count=ncols,
            allow_percent=True,
            field_name="widths",
        )
        column_widths = [ast.Length(value=value, unit=unit) for value, unit in pairs]
    raw_heights = caption.attrs.get("heights")
    if raw_heights:
        pairs = parse_size_list(
            raw_heights,
            expected_count=nrows,
            allow_percent=False,
            field_name="heights",
        )
        row_heights = [ast.Length(value=value, unit=unit) for value, unit in pairs]
    return column_widths, row_heights


def _asset_url(path: str) -> str:
    encoded = quote(path, safe="")
    return f"{_ASSET_ROUTE_PREFIX}/{encoded}"


def _image_aspect_ratio(
    node: ast.Image,
    *,
    metadata: _ImageMetadata | None = None,
) -> str:
    width = _dimension_px(node.width)
    height = _dimension_px(node.height)
    if width is not None and height is not None and width > 0 and height > 0:
        return f"{_format_ratio_number(width)} / {_format_ratio_number(height)}"
    if metadata is not None and metadata.width_px > 0 and metadata.height_px > 0:
        return (
            f"{_format_ratio_number(metadata.width_px)} / "
            f"{_format_ratio_number(metadata.height_px)}"
        )
    return f"{_DEFAULT_IMAGE_ASPECT_WIDTH} / {_DEFAULT_IMAGE_ASPECT_HEIGHT}"


def _dimension_px(spec: ast.Length | None) -> float | None:
    if spec is None or spec.unit == "auto":
        return None
    unit = spec.unit
    value = spec.value
    if unit == "px":
        return value
    if unit == "in":
        return value * _CSS_PX_PER_INCH
    if unit == "cm":
        return value * _CSS_PX_PER_INCH / 2.54
    if unit == "mm":
        return value * _CSS_PX_PER_INCH / 25.4
    if unit == "pt":
        return value * _CSS_PX_PER_INCH / 72
    if unit == "emu":
        return value * _CSS_PX_PER_INCH / _EMU_PER_INCH
    return None


def _format_ratio_number(value: float) -> str:
    return f"{value:.4f}".rstrip("0").rstrip(".")


def _standalone_image(children: list[ast.InlineNode]) -> ast.Image | None:
    meaningful = [
        child
        for child in children
        if not (isinstance(child, ast.Text) and child.text.strip() == "")
    ]
    if len(meaningful) == 1 and isinstance(meaningful[0], ast.Image):
        return meaningful[0]
    return None


def _format_caption(
    config: Config,
    *,
    category: str,
    number: int | str,
    text: str | None,
) -> str:
    spec = _caption_spec(config, category)
    label = _CAPTION_LABELS.get(category, category.capitalize())
    try:
        formatted = spec.format.format(
            category=label,
            number=number,
            text=text or "",
        )
    except (KeyError, IndexError):
        formatted = f"{label} {number}"
        if text:
            formatted = f"{formatted} — {text}"
        return formatted
    if not text:
        formatted = formatted.rstrip()
        for trailing in (" —", " -", "—", "-"):
            if formatted.endswith(trailing):
                formatted = formatted[: -len(trailing)].rstrip()
                break
    return formatted


def _caption_spec(config: Config, category: str) -> CaptionStyle:
    spec = getattr(config.captions, category)
    if isinstance(spec, CaptionStyle):
        return spec
    return CaptionStyle()


def _estimate_image_height(
    node: ast.Image,
    *,
    content_width: Length,
    content_height: Length,
    metadata: _ImageMetadata | None = None,
) -> Length:
    explicit_width = _convert_image_length(node.width, percent_base=content_width)
    explicit_height = _convert_image_length(node.height, percent_base=content_height)
    if explicit_height is not None:
        return explicit_height
    width = explicit_width if explicit_width is not None else content_width
    if metadata is not None:
        if explicit_width is None:
            width = Length(min(int(metadata.native_width), int(content_width)))
        return Length(
            max(
                _MIN_LAYOUT_EMU,
                int(width) * metadata.height_px // metadata.width_px,
            )
        )
    return Length(
        max(
            _MIN_LAYOUT_EMU,
            int(width) * _DEFAULT_IMAGE_ASPECT_HEIGHT // _DEFAULT_IMAGE_ASPECT_WIDTH,
        )
    )


def _convert_image_length(
    spec: ast.Length | None,
    *,
    percent_base: Length,
) -> Length | None:
    if spec is None or spec.unit == "auto":
        return None
    unit = spec.unit
    value = spec.value
    if unit == "%":
        return Length(int(float(percent_base) * value / 100.0))
    if unit == "cm":
        return Length(int(value * 10 * _EMU_PER_MM))
    if unit == "mm":
        return Length(int(value * _EMU_PER_MM))
    if unit == "pt":
        return Pt(value)
    if unit == "in":
        return Length(int(value * _EMU_PER_INCH))
    if unit == "px":
        return Length(int(value * _EMU_PER_INCH / _CSS_PX_PER_INCH))
    if unit == "emu":
        return Length(int(value))
    return None


def _read_filesystem_image_metadata(
    storage: Storage,
    src: str,
) -> _ImageMetadata | None:
    if not isinstance(storage, FilesystemStorage):
        return None
    path = _filesystem_asset_path(storage, src)
    if path is None:
        return None
    try:
        with PILImage.open(path) as image:
            width_px, height_px = image.size
            dpi = image.info.get("dpi")
    except (FileNotFoundError, OSError, UnidentifiedImageError, ValueError):
        return None
    if width_px <= 0 or height_px <= 0:
        return None
    horizontal_dpi, vertical_dpi = _image_dpi(dpi)
    return _ImageMetadata(
        width_px=width_px,
        height_px=height_px,
        horizontal_dpi=horizontal_dpi,
        vertical_dpi=vertical_dpi,
    )


def _filesystem_asset_path(storage: FilesystemStorage, src: str) -> Path | None:
    path = Path(src)
    if path.is_absolute():
        return path
    if storage.base_dir is None:
        return path
    return storage.base_dir / path


def _image_dpi(raw: object) -> tuple[float, float]:
    if not isinstance(raw, tuple) or len(raw) < 2:
        return _DEFAULT_IMAGE_DPI, _DEFAULT_IMAGE_DPI
    horizontal = _positive_float(raw[0]) or _DEFAULT_IMAGE_DPI
    vertical = _positive_float(raw[1]) or _DEFAULT_IMAGE_DPI
    return horizontal, vertical


def _positive_float(value: object) -> float | None:
    if not isinstance(value, str | bytes | bytearray | int | float):
        return None
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    if out <= 0:
        return None
    return out


def _length_to_points(length: Length) -> float:
    return int(length) * 72.0 / _EMU_PER_INCH


def _estimate_table_height(
    node: ast.Table,
    *,
    config: Config,
    content_width: Length,
    row_heights: list[ast.Length] | None,
) -> Length:
    font_size = config.table.font_size or config.font.size
    total = 0
    for row_index, row in enumerate(node.rows):
        if row_heights is not None:
            explicit = row_heights[row_index]
            if explicit.unit != "auto":
                converted = _convert_image_length(explicit, percent_base=content_width)
                if converted is not None:
                    total += int(converted)
                    continue
        row_text = " ".join(
            _collect_inline_text(
                cell.children,
                inline_code=config.paragraph.inline_code,
            )
            for cell in row.cells
        )
        total += int(
            _estimate_text_height(
                row_text,
                font_family=config.font.family,
                font_size=font_size,
                line_spacing=config.font.line_spacing,
                content_width=content_width,
            )
        )
        total += int(Pt(4))
    return Length(max(_MIN_LAYOUT_EMU, total))


def _estimate_listing_height(node: ast.Listing, *, config: Config) -> Length:
    text = node.code.rstrip("\n")
    lines = text.split("\n") if text else [""]
    font_pt = parse_pt(config.listing.font.size)
    line_height = int(
        Pt(
            _calibrated_line_height_pt(config.listing.font.family, font_pt)
            * config.listing.line_spacing
        )
    )
    block_padding = int(Pt(4))
    return Length(max(_MIN_LAYOUT_EMU, len(lines) * line_height + block_padding))


def _listing_lines(code: str) -> list[str]:
    text = code.rstrip("\n")
    return text.split("\n") if text else [""]


def _listing_chunk_text(
    lines: list[str],
    *,
    trailing_newline: bool,
) -> str:
    text = "\n".join(lines)
    if trailing_newline:
        return text + "\n"
    return text


def _estimate_equation_height(
    latex: str,
    *,
    config: Config,
    mathml: str | None,
) -> Length:
    font_pt = parse_pt(config.font.size)
    depth = _mathml_depth(mathml) if mathml is not None else 0
    complexity = max(1, latex.count("\n") + 1)
    base = int(Pt(font_pt * 1.5 * (1 + 0.35 * depth)))
    before = parse_length(config.equation.space_before)
    after = parse_length(config.equation.space_after)
    return Length(max(_MIN_LAYOUT_EMU, base * complexity + int(before) + int(after)))


def _equation_mathml(latex: str) -> tuple[str | None, str | None]:
    try:
        return latex_to_mathml(latex), None
    except EquationError as exc:
        return None, str(exc)


def _mathml_depth(mathml: str | None) -> int:
    if not mathml:
        return 0
    tags = (
        "<mfrac",
        "<msqrt",
        "<mroot",
        "<munderover",
        "<munder",
        "<mover",
        "<msubsup",
        "<msub",
        "<msup",
        "<mtable",
    )
    return sum(mathml.count(tag) for tag in tags)


def _collect_inline_text(
    children: list[ast.InlineNode],
    *,
    inline_code: InlineCode | None = None,
) -> str:
    parts: list[str] = []

    def walk(node: ast.InlineNode) -> None:
        if isinstance(node, ast.Text):
            parts.append(node.text)
            return
        if isinstance(node, ast.InlineCode):
            text = node.code
            if inline_code is not None and inline_code.quotes:
                text = f"«{text}»"
            parts.append(text.replace("-", "‑"))
            return
        if isinstance(node, ast.LineBreak):
            parts.append(" ")
            return
        if isinstance(node, ast.Reference):
            raise ValueError(f"unresolved reference: {node.name}")
            return
        if isinstance(node, ast.Citation):
            raise ValueError(f"unresolved citation: {node.key}")
            return
        if isinstance(node, ast.InlineEquation):
            parts.append(node.latex)
            return
        if isinstance(
            node,
            ast.Link | ast.Emphasis | ast.Strong | ast.Strikethrough | ast.Underline,
        ):
            for child in node.children:
                walk(child)

    for child in children:
        walk(child)
    return "".join(parts).strip()


def _list_item_blocks(
    item: ast.ListItem,
) -> list[tuple[str, list[ast.InlineNode] | ast.List]]:
    blocks: list[tuple[str, list[ast.InlineNode] | ast.List]] = []
    for child in item.children:
        if isinstance(child, ast.Paragraph):
            blocks.append(("para", list(child.children)))
        elif isinstance(child, ast.List):
            blocks.append(("list", child))
        elif isinstance(child, ast.Text):
            blocks.append(("para", [child]))
    return blocks


def _resolve_list_marker_style(node: ast.List) -> str:
    if not node.ordered:
        return "bullet"
    if node.marker_style in _ALPHABETS_BY_STYLE:
        return node.marker_style
    return "arabic"


def _format_list_marker(
    config: Config,
    *,
    marker_style: str,
    counter: int,
    full_path: list[int] | None,
    delimiter: str | None,
    bullet_rel_level: int = 0,
) -> str:
    if marker_style == "bullet":
        # GOST-маппинг DOCX-рендера: L1 — bullet_marker, глубже в режиме
        # native — bullet_nested_format с локальным счётчиком уровня;
        # в режиме inline все уровни — bullet_marker.
        if config.lists.mode == "native" and bullet_rel_level >= 1:
            fmt = config.lists.bullet_nested_format
            try:
                return fmt.format(n=counter)
            except (KeyError, IndexError):
                return fmt
        return config.lists.bullet_marker
    if marker_style in _ALPHABETS_BY_STYLE:
        return _format_alpha_marker(
            config.lists.alphabetic_format,
            counter,
            _ALPHABETS_BY_STYLE[marker_style],
        )
    if full_path:
        number: str | int = _HIERARCHICAL_SEPARATOR.join(str(c) for c in full_path)
    else:
        number = counter
    if delimiter is not None:
        return f"{number}{delimiter}"
    try:
        return config.lists.numbered_format.format(n=number)
    except (KeyError, IndexError):
        return config.lists.numbered_format


def _format_alpha_marker(fmt: str, counter: int, alphabet: str) -> str:
    idx = max(0, counter - 1)
    letter = str(counter) if idx >= len(alphabet) else alphabet[idx]
    try:
        return fmt.format(a=letter, n=counter)
    except (KeyError, IndexError):
        return fmt


def _is_structural_heading(config: Config, node: ast.Heading, raw_text: str) -> bool:
    if node.numbered or node.level != 1:
        return False
    normalized = _normalize_structural_title(raw_text)
    titles = [_normalize_structural_title(title) for title in config.headings.structural_titles]
    return normalized in titles or (
        normalized.startswith("SPISOK ") and any(title.startswith("SPISOK ") for title in titles)
    )


def _normalize_structural_title(text: str) -> str:
    return " ".join(text.upper().strip().rstrip(".").split())
