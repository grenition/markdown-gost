from __future__ import annotations

from collections.abc import Iterator
from itertools import chain
from typing import TYPE_CHECKING, Any, Protocol, cast, runtime_checkable

from docx.document import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Length, Parented

from markdown_gost.render.layout_tracker import LayoutTracker
from markdown_gost.render.numberer import Numberer
from markdown_gost.render.render_index import RenderIndex
from markdown_gost.renderable._oxml import create_element
from markdown_gost.renderable.base import (
    Renderable,
    RenderedInfo,
    RequiresNumbering,
    SubRenderable,
)

if TYPE_CHECKING:
    from markdown_gost.renderable.appendix import Appendix
    from markdown_gost.renderable.heading import Heading

# Bottom margin used for the layout area when the document does not specify one
# explicitly. Matches the legacy renderer; will become configurable in T010+.
_DEFAULT_BOTTOM_MARGIN: Length = Cm(1.86)


@runtime_checkable
class _RenderHook(Protocol):
    """Optional debugger/observer hook called on each rendered element."""

    def add(self, element: Parented, height: Length) -> None: ...
    def after_rendered(self) -> None: ...


def _create_page_field() -> Any:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE \\* MERGEFORMAT")
    return field


class Renderer:
    """Walks a list of :class:`Renderable` objects and writes them into a docx.

    Responsibilities:

    * keep a :class:`LayoutTracker` in sync with the appended elements;
    * detect "this won't fit on the current page" and flush before placing;
    * carry a queue of renderables explicitly deferred to the next page
      (via :class:`SubRenderable.add_to_new_page`);
    * allocate cross-reference numbers for :class:`RequiresNumbering` items;
    * собирать :class:`RenderIndex` за первый проход (T016) — заголовки,
      нумерованные объекты, фактическое число страниц — и вызывать
      :meth:`Renderable.post_process` на втором проходе.
    """

    def __init__(
        self,
        document: Document,
        hook: _RenderHook | None = None,
    ) -> None:
        self._document = document
        self._numberer = Numberer()
        self._hook = hook
        self._index = RenderIndex()
        # Inкрементальный счётчик закладок: уникальный ``w:id`` обязателен,
        # иначе Word/LO молча сольют разные диапазоны в один.
        self._next_bookmark_id = 0

        section = document.sections[0]
        bottom_margin: Length = section.bottom_margin or _DEFAULT_BOTTOM_MARGIN
        page_height: Length = section.page_height or Length(0)
        page_width: Length = section.page_width or Length(0)
        top_margin: Length = section.top_margin or Length(0)
        left_margin: Length = section.left_margin or Length(0)
        right_margin: Length = section.right_margin or Length(0)
        max_height = Length(page_height - top_margin - bottom_margin)
        max_width = Length(page_width - left_margin - right_margin)
        self.layout_tracker = LayoutTracker(max_height, max_width)

        # Page number in the footer. Done unconditionally because it is part of
        # the GOST template baseline; T010+ will make this configurable.
        footer_paragraph = document.sections[0].footer.paragraphs[0]
        footer_paragraph.paragraph_format.first_line_indent = 0
        footer_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        footer_paragraph._p.append(_create_page_field())

        self.previous_rendered: RenderedInfo | None = None
        self._to_new_page: list[Renderable] = []

    @property
    def numberer(self) -> Numberer:
        return self._numberer

    @property
    def index(self) -> RenderIndex:
        """RenderIndex, накопленный за первый проход — для post_process и тестов."""

        return self._index

    def process(self, renderables: list[Renderable]) -> None:
        for renderable in renderables:
            self.render(renderable)

        self._flush_to_new_page()

        # Завершаем первый проход: фиксируем фактическое число страниц.
        # ``layout_tracker.current_state.page`` — текущая «активная» страница;
        # если последний элемент не дотянул до её низа, она всё равно учтена.
        self._index.set_total_pages(self.layout_tracker.current_state.page)

        # Второй проход: даём renderable'ам возможность подставить значения
        # из index (например, total_pages в footer-плейсхолдер). Дефолтная
        # реализация — no-op, поэтому overhead минимален.
        for renderable in renderables:
            renderable.post_process(self._index, self._document)

        if self._hook is not None:
            self._hook.after_rendered()

    def render(self, renderable: Renderable) -> None:
        self._sync_appendix_mode(renderable)
        # Пустая ``numbering_category`` — instance-level opt-out из нумерации
        # (например, ``ast.Listing`` без подписи не получает номер).
        requires_numbering = (
            isinstance(renderable, RequiresNumbering)
            and bool(renderable.numbering_category)
        )
        number: int | None = None
        if requires_numbering:
            assert isinstance(renderable, RequiresNumbering)
            number = self._numberer.get_current_number(renderable.numbering_category) + 1
            renderable.set_number(self._display_number(renderable.numbering_category, number))

        infos: Iterator[RenderedInfo | SubRenderable] = renderable.render(
            self.previous_rendered, self.layout_tracker.current_state
        )

        # Peek the first item: if it cannot fit on the current page, flush
        # then re-invoke render() to give the renderable a fresh layout state.
        try:
            first = next(infos)
            if (
                isinstance(first, RenderedInfo)
                and first.height >= self.layout_tracker.current_state.remaining_page_height
            ):
                self._flush_to_new_page()
                infos = renderable.render(self.previous_rendered, self.layout_tracker.current_state)
            else:
                infos = chain([first], infos)
        except StopIteration:
            return

        first_added_element: Any | None = None
        placed_page: int | None = None
        for info in infos:
            if isinstance(info, SubRenderable):
                if info.add_to_new_page:
                    self._to_new_page.append(info.renderable)
                else:
                    self.render(info.renderable)
            else:
                self._add(info.docx_element, info.height)
                if first_added_element is None:
                    first_added_element = info.docx_element
                    placed_page = self.layout_tracker.current_state.page
                self.previous_rendered = info

        if requires_numbering and number is not None:
            assert isinstance(renderable, RequiresNumbering)
            self._numberer.save_number(renderable.numbering_category, number)

        # Регистрация в RenderIndex — после того, как все физические элементы
        # renderable'а уже вставлены в body; closure-биндинг между записью
        # индекса и закладкой (один общий ``anchor``) удерживается локально.
        self._record_in_index(
            renderable, first_added_element, placed_page, number
        )

    def _flush_to_new_page(self) -> None:
        while self._to_new_page:
            renderable = self._to_new_page.pop(0)
            self._sync_appendix_mode(renderable)
            number: int | None = None
            requires_numbering = (
                isinstance(renderable, RequiresNumbering)
                and bool(renderable.numbering_category)
            )
            if requires_numbering:
                assert isinstance(renderable, RequiresNumbering)
                number = self._numberer.get_current_number(renderable.numbering_category) + 1
                renderable.set_number(self._display_number(renderable.numbering_category, number))
                self._numberer.save_number(renderable.numbering_category, number)
            first_added_element: Any | None = None
            placed_page: int | None = None
            for info in renderable.render(
                self.previous_rendered, self.layout_tracker.current_state
            ):
                if isinstance(info, RenderedInfo):
                    self._add(info.docx_element, info.height)
                    if first_added_element is None:
                        first_added_element = info.docx_element
                        placed_page = self.layout_tracker.current_state.page
                    self.previous_rendered = info
            self._record_in_index(
                renderable, first_added_element, placed_page, number
            )

    # ---- индексация --------------------------------------------------------

    def _record_in_index(
        self,
        renderable: Renderable,
        first_added_element: Any | None,
        placed_page: int | None,
        number: int | None,
    ) -> None:
        """Добавить renderable в :class:`RenderIndex` и повесить закладку.

        Закладка ставится на ``first_added_element`` — первый физически
        размещённый кусок renderable'а (для Heading это сам параграф,
        для Image — параграф с картинкой, для Table/Listing — caption-параграф,
        для Equation — таблица-обёртка). Этого достаточно для гиперссылок
        в DOCX/HTML: переход «к началу объекта» ведёт на правильное место.
        """

        if first_added_element is None or placed_page is None:
            return
        if renderable.identifier:
            self._wrap_with_bookmark(first_added_element, renderable.identifier)

        # Локальный импорт: Heading импортирует базовые классы из
        # ``renderable.base`` — циклически тащить его в ``Renderer`` нет смысла.
        from markdown_gost.renderable.appendix import Appendix as _Appendix
        from markdown_gost.renderable.heading import Heading as _Heading

        if isinstance(renderable, _Appendix):
            self._record_appendix(renderable, first_added_element, placed_page)
            return

        if isinstance(renderable, _Heading):
            self._record_heading(renderable, first_added_element)
            return

        if (
            isinstance(renderable, RequiresNumbering)
            and renderable.numbering_category
            and number is not None
        ):
            self._record_numbered_object(
                renderable, first_added_element, placed_page, number
            )

    def _record_heading(
        self, heading: Heading, element: Any
    ) -> None:
        # ``rendered_page`` выставлен Heading.render() и точнее, чем
        # current_state.page (учитывает page_break_before и порядок add_height).
        page = heading.rendered_page or self.layout_tracker.current_state.page
        entry = self._index.add_heading(
            level=heading.level,
            text=heading.raw_text,
            numbered=heading.is_numbered,
            number=heading.number,
            page=page,
        )
        self._wrap_with_bookmark(element, entry.anchor)

    def _record_appendix(
        self, appendix: Appendix, element: Any, page: int
    ) -> None:
        entry = self._index.add_heading(
            level=1,
            text=appendix.index_text,
            numbered=False,
            number=None,
            page=page,
        )
        self._wrap_with_bookmark(element, entry.anchor)

    def _record_numbered_object(
        self,
        renderable: Renderable,
        element: Any,
        page: int,
        number: int,
    ) -> None:
        category = cast(RequiresNumbering, renderable).numbering_category
        caption: str | None = getattr(renderable, "caption_text", None)
        entry = self._index.add_object(
            category=category,
            number=number,
            number_text=(
                self._numberer.format_number(category, number)
                if self._numberer.current_appendix is not None
                else None
            ),
            caption=caption,
            page=page,
        )
        self._wrap_with_bookmark(element, entry.anchor)

    def _sync_appendix_mode(self, renderable: Renderable) -> None:
        from markdown_gost.renderable.appendix import Appendix as _Appendix
        from markdown_gost.renderable.appendix import AppendixEnd

        if isinstance(renderable, _Appendix):
            self._numberer.start_appendix(renderable.letter)
        elif isinstance(renderable, AppendixEnd):
            self._numberer.end_appendix()

    def _display_number(self, category: str, number: int) -> int | str:
        if self._numberer.current_appendix is None:
            return number
        return self._numberer.format_number(category, number)

    def _wrap_with_bookmark(self, element: Any, anchor: str) -> None:
        """Вставить ``<w:bookmarkStart>`` перед ``element`` и ``<w:bookmarkEnd>``
        после него (на уровне ``<w:body>``).

        Bookmark на уровне body — самый совместимый вариант для гиперссылок:
        работает и для параграфов, и для таблиц, не ломает inline-структуру
        внутри runs. Word/LO принимают такие закладки как «начало диапазона»
        и кликабельная ссылка прыгает в нужное место.
        """

        xml_element = getattr(element, "_element", element)
        parent = xml_element.getparent()
        if parent is None:
            return

        bm_id = self._next_bookmark_id
        self._next_bookmark_id += 1

        start = create_element(
            "w:bookmarkStart", {"w:id": str(bm_id), "w:name": anchor}
        )
        end = create_element("w:bookmarkEnd", {"w:id": str(bm_id)})

        idx = list(parent).index(xml_element)
        parent.insert(idx, start)
        # После вставки start индекс ``element`` сдвинулся на 1 → end после него.
        parent.insert(idx + 2, end)

    def _add(self, element: Parented, height: Length) -> None:
        # python-docx exposes the body as a private attribute; the public API
        # has no clean way to append a pre-built Parented at the end of the body
        body = cast(Any, self._document)._body
        body._element.append(cast(Any, element)._element)
        self.layout_tracker.add_height(height)

        if self._hook is not None:
            self._hook.add(element, height)
