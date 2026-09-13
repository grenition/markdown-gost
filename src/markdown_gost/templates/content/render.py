"""Оглавление документа с кликабельными ссылками.

Каждая запись TOC рендерится отдельным DOCX-параграфом с правым tab stop и
dot leader для номера страницы. Текст записи обёрнут в ``<w:hyperlink
w:anchor="…">`` с явным ``<w:color w:val="auto"/>`` и без подчёркивания —
клик в Word/LO ведёт к закладке заголовка, текст остаётся чёрным. После TOC
принудительный разрыв страницы.

Заголовок «СОДЕРЖАНИЕ» рендерится самим шаблоном (параметр ``title``) —
стилизуется как structural heading, но **не проходит через ``Numberer``** и
поэтому не занимает слот в нумерации (``# Введение`` после TOC остаётся
``"1"``, а не ``"2"``). Текст заголовка можно изменить, пустая строка
полностью отключает его. Этот заголовок (как и любые другие, попавшие на
страницу самого TOC до его параграфа) исключается из списка TOC —
иначе он бы ссылался сам на себя и засорял оглавление.

Нумерация в TOC берётся из ``RenderIndex.HeadingEntry.number``. Префикс
рендерится только если ``numbered=True``.
"""

from __future__ import annotations

from collections.abc import Generator
from copy import copy
from typing import Any, cast

from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm, Length
from docx.text.paragraph import Paragraph as DocxParagraph

from markdown_gost.config.schema import Config
from markdown_gost.render.document_factory import STRUCTURAL_HEADING_STYLE
from markdown_gost.render.layout_tracker import LayoutState
from markdown_gost.render.paragraph_sizer import ParagraphSizer
from markdown_gost.render.render_index import HeadingEntry, RenderIndex
from markdown_gost.renderable._oxml import create_element
from markdown_gost.renderable.base import Renderable, RenderedInfo, SubRenderable
from markdown_gost.renderable.page_break import PageBreak

_TOC_LEVEL_INDENT = Cm(0.75)
_TOC_HANGING_INDENT = Cm(0.75)


class Content(Renderable):
    """Оглавление, заполняемое во второй проход (T016).

    Параметры:

    * ``title`` — заголовок перед TOC (по умолчанию «СОДЕРЖАНИЕ», пустая
      строка отключает); рендерится как structural heading, но в отличие от
      обычного заголовка не консумит слот в нумерации;
    * ``depth`` — макс. уровень включаемых заголовков (1..6, по умолчанию 6);
    * ``dot_leader`` — точечный leader между названием и страницей.
    """

    def __init__(
        self,
        parent: Any,
        config: Config,
        params: dict[str, Any],
    ) -> None:
        self._parent = parent
        self._config = config
        self._title: str = str(params.get("title", "СОДЕРЖАНИЕ"))
        self._depth: int = int(params.get("depth", 6))
        self._dot_leader: bool = bool(params.get("dot_leader", True))
        self._toc_page: int = 1

        self._docx_paragraph = DocxParagraph(create_element("w:p"), parent)
        self._docx_paragraph.style = "Normal"
        pf = self._docx_paragraph.paragraph_format
        pf.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        pf.first_line_indent = 0

    # ---- render -----------------------------------------------------------

    def render(
        self,
        previous_rendered: RenderedInfo | None,
        layout_state: LayoutState,
    ) -> Generator[RenderedInfo | SubRenderable]:
        if self._title:
            yield from self._render_title(previous_rendered, layout_state)
            previous_rendered = None
        # Запоминаем номер страницы, на которой окажется параграф TOC. По
        # ней в post_process отфильтруем заголовки, попавшие на ту же
        # страницу до TOC (включая собственный ``title``) — они не должны
        # появляться в самом оглавлении.
        self._toc_page = layout_state.page
        yield RenderedInfo(self._docx_paragraph, Length(0))
        yield from PageBreak(self._parent).render(None, copy(layout_state))

    def _render_title(
        self,
        previous_rendered: RenderedInfo | None,
        layout_state: LayoutState,
    ) -> Generator[RenderedInfo | SubRenderable]:
        """Эмитит заголовок TOC как structural heading без участия ``Numberer``.

        Минусом обхода обычного ``Heading`` renderable является дублирование
        page-break-логики, плюс — заголовок не занимает слот в нумерации.
        """

        title_p = DocxParagraph(create_element("w:p"), self._parent)
        title_p.style = STRUCTURAL_HEADING_STYLE
        title_p.add_run(self._title)

        previous_docx: DocxParagraph | None = None
        if previous_rendered is not None and isinstance(
            previous_rendered.docx_element, DocxParagraph
        ):
            previous_docx = previous_rendered.docx_element

        height_data = ParagraphSizer(
            title_p,
            previous_docx,
            layout_state.max_width,
        ).calculate_height()

        # Structural title — с новой страницы, если на текущей уже что-то лежит.
        remaining = layout_state.remaining_page_height
        page_break = (
            self._config.headings.structural.page_break_before
            and layout_state.current_page_height > 0
        )
        if page_break:
            title_p.paragraph_format.page_break_before = True

        if layout_state.current_page_height == 0 and layout_state.page != 1:
            height_data.before = Length(0)

        height = Length(int(height_data.full))
        if page_break:
            height = Length(int(height) + int(remaining))

        layout_state.add_height(height)
        yield RenderedInfo(title_p, height)

    # ---- post_process -----------------------------------------------------

    def post_process(self, index: RenderIndex, document: Any) -> None:
        section = document.sections[0]
        content_width = (
            section.page_width - section.left_margin - section.right_margin
        )

        # Исключаем заголовки, попавшие на страницу самого TOC до него
        # (например, ненумерованное «СОДЕРЖАНИЕ») — они страница ≤ TOC.
        toc_entries = [
            h
            for h in index.headings
            if h.level <= self._depth and h.page > self._toc_page
        ]
        toc_paragraphs = self._ensure_toc_paragraphs(len(toc_entries))
        for paragraph, entry in zip(toc_paragraphs, toc_entries, strict=True):
            self._render_entry_paragraph(paragraph, entry, content_width)

    def _ensure_toc_paragraphs(self, count: int) -> list[DocxParagraph]:
        """Вернуть ``count`` соседних TOC-параграфов, начиная с placeholder."""

        self._clear_paragraph(self._docx_paragraph)
        if count == 0:
            return []

        paragraphs = [self._docx_paragraph]
        previous = self._docx_paragraph._element
        for _ in range(count - 1):
            element = create_element("w:p")
            previous.addnext(element)
            paragraph = DocxParagraph(element, self._parent)
            paragraphs.append(paragraph)
            previous = element
        return paragraphs

    def _render_entry_paragraph(
        self,
        paragraph: DocxParagraph,
        entry: HeadingEntry,
        content_width: Length,
    ) -> None:
        paragraph.style = "Normal"
        pf = paragraph.paragraph_format
        pf.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        level_offset = Length(max(entry.level - 1, 0) * int(_TOC_LEVEL_INDENT))
        if entry.numbered and entry.number:
            pf.left_indent = Length(int(level_offset) + int(_TOC_HANGING_INDENT))
            pf.first_line_indent = Length(-int(_TOC_HANGING_INDENT))
            paragraph.add_run(f"{entry.number} ")
        else:
            pf.left_indent = level_offset
            pf.first_line_indent = 0

        leader_right = (
            WD_TAB_LEADER.DOTS if self._dot_leader else WD_TAB_LEADER.SPACES
        )
        pf.tab_stops.add_tab_stop(
            content_width,
            alignment=WD_TAB_ALIGNMENT.RIGHT,
            leader=leader_right,
        )

        self._add_link(paragraph, entry.text, entry.anchor)
        paragraph.add_run(f"\t{entry.page}")

    def _clear_paragraph(self, paragraph: DocxParagraph) -> None:
        cast(Any, paragraph._p).clear_content()

    def _add_link(
        self, paragraph: DocxParagraph, text: str, anchor: str
    ) -> None:
        """Добавить кликабельную ссылку на закладку ``anchor`` чёрным цветом.

        Стиль ``Hyperlink`` сделал бы её синей с подчёркиванием — мы хотим
        обычный текст, поэтому вместо стиля выставляем явный
        ``<w:color w:val="auto"/>`` и подчёркивание не включаем.
        """

        hyperlink = create_element("w:hyperlink", {"w:anchor": anchor})
        hyperlink.append(
            create_element(
                "w:r",
                [
                    create_element(
                        "w:rPr",
                        [create_element("w:color", {"w:val": "auto"})],
                    ),
                    create_element(
                        "w:t",
                        {"xml:space": "preserve"},
                        text,
                    ),
                ],
            )
        )
        paragraph._p.append(hyperlink)


def render(params: dict[str, Any], parent: Any) -> Renderable:
    """Точка входа шаблона; ``parent`` — :class:`TemplateContext`."""

    return Content(parent.document, parent.config, params)
