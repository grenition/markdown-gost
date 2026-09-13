"""Титульный лист по типовой российской университетской форме — порт легаси-рендера 1:1.

Структура (см. legacy ``md2gost/renderable/title_page.py``): логотип →
«МИНОБРНАУКИ РОССИИ» → полное название федерального учреждения курсивом →
полное наименование вуза → сокращённое наименование вуза →
горизонтальная линия → институт → кафедра → 3 пустых → заголовок работы 16pt
bold → пустой → ``по дисциплине «...»`` → 6 пустых → блок «Выполнил/группа/
Проверил» → город+год, прижатые к низу страницы через ``framePr``.

Все строки идут одной серией параграфов с ``height=0``; высоту страницы
выбирает завершающий пустой параграф — это гарантирует, что основной текст
стартует со следующей страницы.
"""

from __future__ import annotations

import io
import urllib.request
from collections.abc import Generator
from pathlib import Path
from typing import Any

from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm, Length, Pt
from docx.text.paragraph import Paragraph as DocxParagraph

from markdown_gost.config.schema import Config
from markdown_gost.render.layout_tracker import LayoutState
from markdown_gost.renderable._oxml import create_element
from markdown_gost.renderable.base import Renderable, RenderedInfo, SubRenderable

_LOGO_MAX_BYTES = 5 * 1024 * 1024
_LOGO_TIMEOUT_SECONDS = 10.0
_LOGO_WIDTH = Cm(2.5)
_FONT_NAME = "Times New Roman"
_DEFAULT_FONT_SIZE: Length = Pt(14)
_ZERO_LENGTH: Length = Length(0)


def _load_logo(reference: object) -> bytes | None:
    """Load an optional emblem image: local path or http(s) URL.

    A broken explicit reference is an error (it must not silently vanish);
    an absent/empty ``logo`` simply renders no emblem.
    """

    if not reference or not isinstance(reference, str):
        return None
    ref = reference.strip()
    if not ref:
        return None
    if ref.startswith(("http://", "https://")):
        with urllib.request.urlopen(ref, timeout=_LOGO_TIMEOUT_SECONDS) as response:
            data = response.read(_LOGO_MAX_BYTES + 1)
    else:
        path = Path(ref)
        if not path.is_file():
            raise ValueError(f"logo not found: {ref!r}")
        data = path.read_bytes()
    if len(data) > _LOGO_MAX_BYTES:
        raise ValueError(f"logo exceeds {_LOGO_MAX_BYTES} bytes: {ref!r}")
    return data


def _add_centered_paragraph(
    parent: Any,
    text: str,
    *,
    font_size: Length = _DEFAULT_FONT_SIZE,
    bold: bool = False,
    italic: bool = False,
    space_before: Length = _ZERO_LENGTH,
    space_after: Length = _ZERO_LENGTH,
) -> DocxParagraph:
    p = DocxParagraph(create_element("w:p"), parent)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.line_spacing = 1
    if text:
        run = p.add_run(text)
        run.font.name = _FONT_NAME
        run.font.size = font_size
        run.font.bold = bold
        run.font.italic = italic
    return p


def _add_empty_paragraph(
    parent: Any,
    *,
    space_before: Length = _ZERO_LENGTH,
    space_after: Length = _ZERO_LENGTH,
) -> DocxParagraph:
    return _add_centered_paragraph(
        parent,
        "",
        space_before=space_before,
        space_after=space_after,
    )


class Titlepage(Renderable):
    """Титульный лист — рендер по типовой университетской легаси-форме (1:1)."""

    def __init__(self, parent: Any, config: Config, params: dict[str, Any]) -> None:
        self._parent = parent
        self._config = config
        self._params = params

    def render(
        self,
        previous_rendered: RenderedInfo | None,
        layout_state: LayoutState,
    ) -> Generator[RenderedInfo | SubRenderable]:
        parent = self._parent
        data = self._params
        page_height = layout_state.max_height
        elements: list[DocxParagraph] = []

        # На странице с титульником номер страницы не должен отображаться.
        # Включаем «Different first page» для секции — первый footer останется
        # пустым, основной (с PAGE-полем, выставленным Renderer'ом) показывается
        # начиная со второй страницы.
        parent.sections[0].different_first_page_header_footer = True

        # --- Logo (optional emblem: ``logo`` param — path or http(s) URL) ---
        logo_p = DocxParagraph(create_element("w:p"), parent)
        logo_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        logo_p.paragraph_format.space_before = 0
        logo_p.paragraph_format.space_after = 0
        logo_p.paragraph_format.first_line_indent = 0
        logo_p.paragraph_format.line_spacing = 1
        run = logo_p.add_run()
        logo_data = _load_logo(data.get("logo"))
        if logo_data is not None:
            run.add_picture(io.BytesIO(logo_data), width=_LOGO_WIDTH)
        elements.append(logo_p)

        # --- МИНОБРНАУКИ РОССИИ ---
        elements.append(_add_centered_paragraph(parent, "МИНОБРНАУКИ РОССИИ"))

        # --- Полное название федерального учреждения (курсив, 11pt) ---
        elements.append(
            _add_centered_paragraph(
                parent,
                "Федеральное государственное бюджетное образовательное учреждение\n"
                "высшего образования",
                font_size=Pt(11),
                italic=True,
            )
        )

        # --- Полное наименование вуза (bold) ---
        if data.get("university_full"):
            elements.append(
                _add_centered_paragraph(
                    parent,
                    data["university_full"],
                    bold=True,
                )
            )

        # --- Сокращённое наименование вуза (16pt bold) ---
        if data.get("university_short"):
            elements.append(
                _add_centered_paragraph(
                    parent, data["university_short"], font_size=Pt(16), bold=True
                )
            )

        # --- Горизонтальная линия (нижняя граница пустого параграфа) ---
        hr_p = _add_empty_paragraph(parent, space_after=Pt(6))
        ppr = hr_p._p.get_or_add_pPr()
        pbdr = create_element(
            "w:pBdr",
            [
                create_element(
                    "w:bottom",
                    {
                        "w:val": "single",
                        "w:sz": "4",
                        "w:space": "1",
                        "w:color": "auto",
                    },
                )
            ],
        )
        ppr.append(pbdr)
        elements.append(hr_p)

        # --- Институт ---
        if data.get("institute"):
            elements.append(
                _add_centered_paragraph(
                    parent, data["institute"], space_before=Pt(6)
                )
            )

        # --- Кафедра ---
        if data.get("department"):
            elements.append(_add_empty_paragraph(parent))
            elements.append(_add_centered_paragraph(parent, data["department"]))

        # --- Spacers перед заголовком ---
        for _ in range(3):
            elements.append(_add_empty_paragraph(parent))

        # --- Заголовок работы ---
        if data.get("title"):
            elements.append(
                _add_centered_paragraph(
                    parent, data["title"], font_size=Pt(16), bold=True
                )
            )

        # --- Spacer ---
        elements.append(_add_empty_paragraph(parent))

        # --- Дисциплина ---
        if data.get("subject"):
            elements.append(
                _add_centered_paragraph(
                    parent, f"по дисциплине «{data['subject']}»"
                )
            )

        if data.get("topic"):
            elements.append(_add_centered_paragraph(parent, "на тему"))
            elements.append(
                _add_centered_paragraph(parent, f"«{data['topic']}»", bold=True)
            )

        # --- Spacers перед таблицей авторов ---
        for _ in range(6):
            elements.append(_add_empty_paragraph(parent))

        # --- Блок авторов / рецензента ---
        elements.extend(self._build_author_block(parent, data))

        # --- Footer (город + год), прижатый к низу страницы ---
        city = (data.get("city") or "").strip()
        year = (data.get("year") or "").strip()
        footer_text = " ".join(part for part in (city, year) if part)
        if footer_text:
            sect = parent.sections[0]
            page_h_twips = int(sect.page_height / 635)
            bottom_twips = int(sect.bottom_margin / 635)
            line_twips = int(Pt(14) / 635 * 2.5)
            y_twips = page_h_twips - bottom_twips - line_twips
            content_w_twips = int(
                (sect.page_width - sect.left_margin - sect.right_margin) / 635
            )

            footer_p = _add_centered_paragraph(parent, footer_text)
            ppr = footer_p._p.get_or_add_pPr()
            ppr.append(
                create_element(
                    "w:framePr",
                    {
                        "w:w": str(content_w_twips),
                        "w:hSpace": "0",
                        "w:vAnchor": "page",
                        "w:hAnchor": "margin",
                        "w:x": "0",
                        "w:y": str(y_twips),
                        "w:wrap": "notBeside",
                    },
                )
            )
            elements.append(footer_p)

        # Все элементы кроме последнего идут с height=0; последний параграф
        # содержит физический разрыв страницы (``w:br w:type="page"``) и
        # «съедает» оставшуюся высоту страницы в layout-трекере, чтобы
        # основной текст начался со следующей.
        for el in elements:
            yield RenderedInfo(el, Length(0))
        page_break_p = _add_empty_paragraph(parent)
        page_break_p._p.append(
            create_element("w:r", [create_element("w:br", {"w:type": "page"})])
        )
        yield RenderedInfo(page_break_p, page_height)

    def _build_author_block(
        self, parent: Any, data: dict[str, Any]
    ) -> list[DocxParagraph]:
        """Блок «Выполнил/Проверил».

        ``authors`` и ``reviewers`` — списки **групп** вида
        ``{label: str, names: list[str]}``. Подпись ``label`` задаётся
        пользователем целиком (например, ``"Студент группы ИКБО-20-23"``
        или ``"Старший преподаватель"``) и пишется слева у первого имени
        группы; остальные имена идут справа без префикса. Между группами
        вставляется небольшой вертикальный отступ.
        """

        paragraphs: list[DocxParagraph] = []
        sect = parent.sections[0]
        content_width = sect.page_width - sect.left_margin - sect.right_margin
        right_tab_pos = str(int(content_width / 635))

        def _bold_label(text: str) -> DocxParagraph:
            p = DocxParagraph(create_element("w:p"), parent)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.space_before = 0
            p.paragraph_format.space_after = 0
            p.paragraph_format.line_spacing = 1
            run = p.add_run(text)
            run.font.name = _FONT_NAME
            run.font.size = Pt(14)
            run.font.bold = True
            return p

        def _label_with_name(label: str, name: str) -> DocxParagraph:
            """Параграф ``label\\t<name>`` — name прижат правым tab-стопом."""

            p = DocxParagraph(create_element("w:p"), parent)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.space_before = 0
            p.paragraph_format.space_after = 0
            p.paragraph_format.line_spacing = 1
            if label:
                run_label = p.add_run(label)
                run_label.font.name = _FONT_NAME
                run_label.font.size = Pt(14)
            tab_run = p.add_run("\t")
            tab_run.font.size = Pt(14)
            run_name = p.add_run(name)
            run_name.font.name = _FONT_NAME
            run_name.font.size = Pt(14)
            tabs = create_element(
                "w:tabs",
                [
                    create_element(
                        "w:tab",
                        {"w:val": "right", "w:pos": right_tab_pos},
                    )
                ],
            )
            p._p.get_or_add_pPr().append(tabs)
            return p

        def _right_aligned_name(name: str) -> DocxParagraph:
            """Имя справа без подписи слева — для дополнительных имён в группе."""

            p = DocxParagraph(create_element("w:p"), parent)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.space_before = 0
            p.paragraph_format.space_after = 0
            p.paragraph_format.line_spacing = 1
            run = p.add_run(name)
            run.font.name = _FONT_NAME
            run.font.size = Pt(14)
            return p

        def _normalize_groups(raw: Any) -> list[tuple[str, list[str]]]:
            """``[{label, names}]`` → ``[(label, [names])]``; пустые группы выкинуты."""

            result: list[tuple[str, list[str]]] = []
            for item in raw or []:
                if not isinstance(item, dict):
                    continue
                label = str(item.get("label") or "")
                names_raw = item.get("names") or []
                names = [str(n) for n in names_raw if n]
                if names:
                    result.append((label, names))
            return result

        def _render_section(title: str, groups: list[tuple[str, list[str]]]) -> None:
            if not groups:
                return
            label_p = _bold_label(f"{title}:")
            if paragraphs:
                # Между секциями «Выполнил» и «Проверил» — увеличенный отступ.
                label_p.paragraph_format.space_before = Pt(12)
            paragraphs.append(label_p)
            for idx, (group_label, names) in enumerate(groups):
                first_p = _label_with_name(group_label, names[0])
                if idx > 0:
                    # Между группами в одной секции — небольшой отступ.
                    first_p.paragraph_format.space_before = Pt(6)
                paragraphs.append(first_p)
                for name in names[1:]:
                    paragraphs.append(_right_aligned_name(name))

        _render_section(
            data.get("authors_title") or "Выполнил",
            _normalize_groups(data.get("authors")),
        )
        _render_section(
            data.get("reviewer_title") or "Проверил",
            _normalize_groups(data.get("reviewers")),
        )

        return paragraphs


def render(params: dict[str, Any], parent: Any) -> Renderable:
    """Точка входа шаблона; ``parent`` — :class:`TemplateContext`."""

    return Titlepage(parent.document, parent.config, params)
