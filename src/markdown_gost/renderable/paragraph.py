"""Параграф из AST → docx-параграф с измерением высоты для page-break."""

from __future__ import annotations

import logging
from collections.abc import Generator
from io import BytesIO
from typing import Any

from docx.enum.text import WD_UNDERLINE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Length, Pt, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.text.run import Run as DocxRun

from markdown_gost.config.schema import Config
from markdown_gost.config.units import parse_pt
from markdown_gost.core.ast import nodes as ast
from markdown_gost.image_placeholder import placeholder_copy, placeholder_png
from markdown_gost.render.bibliography import BibliographyIndex
from markdown_gost.render.layout_tracker import LayoutState
from markdown_gost.render.paragraph_sizer import ParagraphSizer
from markdown_gost.renderable.base import Renderable, RenderedInfo, SubRenderable
from markdown_gost.storage.base import Storage, StorageError

from ._oxml import create_element

_log = logging.getLogger(__name__)


class Paragraph(Renderable):
    """Обычный текстовый параграф со встроенными inline-элементами."""

    def __init__(
        self,
        parent: Any,
        config: Config,
        style: str = "Normal",
        *,
        storage: Storage | None = None,
        bibliography: BibliographyIndex | None = None,
    ) -> None:
        self._parent = parent
        self._config = config
        self._storage = storage
        self._bibliography = bibliography
        self._docx_paragraph = DocxParagraph(create_element("w:p"), parent)
        self._docx_paragraph.style = style

    @property
    def docx_paragraph(self) -> DocxParagraph:
        return self._docx_paragraph

    @property
    def page_break_before(self) -> bool:
        return bool(self._docx_paragraph.paragraph_format.page_break_before)

    @page_break_before.setter
    def page_break_before(self, value: bool) -> None:
        self._docx_paragraph.paragraph_format.page_break_before = value

    # ---- inline composition -------------------------------------------------

    def add_inline_nodes(self, children: list[ast.InlineNode]) -> None:
        """Преобразовать AST-инлайны в docx runs."""
        self._add_inline(
            self._docx_paragraph,
            children,
            bold=False,
            italic=False,
            strike=False,
            underline=False,
        )

    def _add_inline(
        self,
        host: Any,
        children: list[ast.InlineNode],
        *,
        bold: bool,
        italic: bool,
        strike: bool,
        underline: bool,
    ) -> None:
        for child in children:
            self._add_inline_node(
                host, child, bold=bold, italic=italic, strike=strike, underline=underline
            )

    def _add_inline_node(
        self,
        host: Any,
        node: ast.InlineNode,
        *,
        bold: bool,
        italic: bool,
        strike: bool,
        underline: bool,
    ) -> None:
        if isinstance(node, ast.Text):
            self._add_run(
                host, node.text, bold=bold, italic=italic, strike=strike, underline=underline
            )
            return
        if isinstance(node, ast.Strong):
            self._add_inline(
                host, node.children, bold=True, italic=italic, strike=strike, underline=underline
            )
            return
        if isinstance(node, ast.Emphasis):
            self._add_inline(
                host, node.children, bold=bold, italic=True, strike=strike, underline=underline
            )
            return
        if isinstance(node, ast.Strikethrough):
            self._add_inline(
                host, node.children, bold=bold, italic=italic, strike=True, underline=underline
            )
            return
        if isinstance(node, ast.Underline):
            self._add_inline(
                host, node.children, bold=bold, italic=italic, strike=strike, underline=True
            )
            return
        if isinstance(node, ast.InlineCode):
            self._add_inline_code(
                host, node, bold=bold, italic=italic, strike=strike, underline=underline
            )
            return
        if isinstance(node, ast.LineBreak):
            # Soft и hard переносы внутри параграфа сворачиваем в пробел —
            # перенос принадлежит layout'у, а не семантике markdown-параграфа.
            self._add_run(host, " ", bold=bold, italic=italic, strike=strike, underline=underline)
            return
        if isinstance(node, ast.InlineEquation):
            self._add_inline_equation(host, node)
            return
        if isinstance(node, ast.Link):
            self._add_link(node, bold=bold, italic=italic, strike=strike, underline=underline)
            return
        if isinstance(node, ast.Reference):
            raise ValueError(f"unresolved reference: {node.name}")
        if isinstance(node, ast.Citation):
            if self._bibliography is None:
                raise ValueError("citation resolver is not configured")
            self._add_run(
                host,
                self._bibliography.format_citation(node.key),
                bold=bold,
                italic=italic,
                strike=strike,
                underline=underline,
            )
            return
        if isinstance(node, ast.Image):
            self._add_inline_image(host, node)
            return
        # Неподдерживаемый inline — текстовая заглушка с типом.
        type_name = type(node).__name__
        self._add_run(
            host,
            f"<{type_name}?>",
            bold=bold,
            italic=italic,
            strike=strike,
            underline=underline,
        )

    def _add_run(
        self,
        host: Any,
        text: str,
        *,
        bold: bool,
        italic: bool,
        strike: bool,
        underline: bool = False,
        font_name: str | None = None,
        font_size: Length | None = None,
        style: str | None = None,
        split_hyphens: bool = True,
    ) -> None:
        if not text:
            return
        if not split_hyphens:
            # Verbatim-режим (inline-код): дефисы остаются как есть.
            self._make_run(
                host,
                text,
                bold=bold,
                italic=italic,
                strike=strike,
                underline=underline,
                font_name=font_name,
                font_size=font_size,
                style=style,
            )
            return
        # Разбиваем по дефисам и заменяем их на noBreakHyphen — поведение из
        # старого рендера: «слова-через-дефис» не должны рваться.
        parts = text.split("-")
        for i, part in enumerate(parts):
            if part:
                self._make_run(
                    host,
                    part,
                    bold=bold,
                    italic=italic,
                    strike=strike,
                    underline=underline,
                    font_name=font_name,
                    font_size=font_size,
                    style=style,
                )
            if i != len(parts) - 1:
                hyphen_run = self._make_run(
                    host,
                    "",
                    bold=bold,
                    italic=italic,
                    strike=strike,
                    underline=underline,
                    font_name=font_name,
                    font_size=font_size,
                    style=style,
                )
                if hyphen_run is not None:
                    hyphen_run._element.append(create_element("w:noBreakHyphen"))

    def _make_run(
        self,
        host: Any,
        text: str,
        *,
        bold: bool,
        italic: bool,
        strike: bool,
        underline: bool = False,
        font_name: str | None = None,
        font_size: Length | None = None,
        style: str | None = None,
    ) -> DocxRun | None:
        if isinstance(host, DocxParagraph):
            run = host.add_run(text)
        else:
            run = DocxRun(create_element("w:r"), self._docx_paragraph)
            host.append(run._element)
            run.text = text
        if style is not None:
            run.style = style
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if strike:
            run.font.strike = True
        if underline:
            run.font.underline = WD_UNDERLINE.SINGLE
        if font_name is not None:
            run.font.name = font_name
        if font_size is not None:
            run.font.size = font_size
        return run

    def _add_inline_code(
        self,
        host: Any,
        node: ast.InlineCode,
        *,
        bold: bool,
        italic: bool,
        strike: bool,
        underline: bool = False,
    ) -> None:
        spec = self._config.paragraph.inline_code
        text = node.code
        if spec.quotes:
            text = f"«{text}»"
        # Inline-код — verbatim, но дефисы не должны рвать строку при переносе.
        # Заменяем U+002D на U+2011 (NON-BREAKING HYPHEN): визуально та же чёрточка,
        # но без точки разрыва. Это сохраняет идентификаторы вида `kebab-case-id`
        # цельными и не порождает артефактов вроде `- - -` для `---`.
        text = text.replace("-", "‑")
        font_size = Pt(parse_pt(spec.size))
        self._add_run(
            host,
            text,
            bold=bold,
            italic=italic or spec.italic,
            strike=strike,
            underline=underline,
            font_name=spec.font,
            font_size=font_size,
            split_hyphens=False,
        )

    def _add_link(
        self,
        node: ast.Link,
        *,
        bold: bool,
        italic: bool,
        strike: bool,
        underline: bool = False,
    ) -> None:
        try:
            if node.url.startswith("#"):
                attrs = {"w:anchor": node.url[1:]}
            else:
                r_id = self._docx_paragraph.part.relate_to(
                    node.url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
                )
                attrs = {"r:id": r_id}
        except Exception:
            # Если parent не подключён к docx-части — fallback в простой run.
            self._add_inline(
                self._docx_paragraph,
                node.children,
                bold=bold,
                italic=italic,
                strike=strike,
                underline=underline,
            )
            return
        hyperlink = create_element("w:hyperlink", attrs)
        self._docx_paragraph._p.append(hyperlink)
        # Применяем стиль "Hyperlink" к runs внутри — это даёт синий + подчёркивание
        # и помогает LibreOffice корректно распознать ссылку при экспорте в PDF.
        self._add_link_inline(
            hyperlink,
            node.children,
            bold=bold,
            italic=italic,
            strike=strike,
            underline=underline,
        )

    def _add_link_inline(
        self,
        host: Any,
        children: list[ast.InlineNode],
        *,
        bold: bool,
        italic: bool,
        strike: bool,
        underline: bool = False,
    ) -> None:
        for child in children:
            if isinstance(child, ast.Text):
                self._add_run(
                    host,
                    child.text,
                    bold=bold,
                    italic=italic,
                    strike=strike,
                    underline=underline,
                    style="Hyperlink",
                )
            else:
                # Вложенные форматирования внутри ссылки рендерим обычным путём.
                self._add_inline_node(
                    host,
                    child,
                    bold=bold,
                    italic=italic,
                    strike=strike,
                    underline=underline,
                )

    def _add_inline_equation(self, host: Any, node: ast.InlineEquation) -> None:
        """Inline-формула ``$...$`` → встроенный ``<m:oMath>`` в run (T015).

        OMML вставляется напрямую в host (``<w:p>`` или OXML hyperlink-обёртка).
        ``<w:rPr>`` math-runs прописывается под текущий ``font.size``/family —
        в Word'е это даёт корректный размер; в LibreOffice math-runs могут
        рендериться меньше из-за известного бага LO-импорта OMML, но формула
        остаётся живой математикой (редактируется в Word).
        """

        from markdown_gost.render.latex_math import (
            EquationError,
            apply_run_props,
            latex_to_omml,
        )

        try:
            omml = latex_to_omml(node.latex)
        except EquationError as exc:
            _log.warning(
                "invalid inline LaTeX equation %r: %s — rendering red placeholder",
                node.latex,
                exc,
            )
            run = self._make_run(
                host,
                f"[invalid: {node.latex}]",
                bold=False,
                italic=True,
                strike=False,
                underline=False,
            )
            if run is not None:
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            return

        apply_run_props(
            omml,
            font_pt=parse_pt(self._config.font.size),
            font_name=self._config.font.family,
        )

        # ``<m:oMath>`` сам себе run-уровень; его можно положить прямо в
        # ``<w:p>`` или в OXML hyperlink-обёртку.
        if isinstance(host, DocxParagraph):
            host._p.append(omml)
        else:
            host.append(omml)

    def _add_inline_image(self, host: Any, node: ast.Image) -> None:
        """Inline-картинка ``![](...)`` внутри текста — без подписи."""

        if self._storage is None:
            self._add_run(
                host,
                f"[{node.alt or node.src}]",
                bold=False,
                italic=True,
                strike=False,
                underline=False,
            )
            return
        try:
            data = self._storage.fetch(node.src)
        except StorageError as exc:
            _log.warning("inline image %r unavailable: %s", node.src, exc)
            self._add_inline_image_placeholder(host, node.alt)
            return
        run = self._make_run(
            host,
            "",
            bold=False,
            italic=False,
            strike=False,
            underline=False,
        )
        if run is None:
            return
        try:
            run.add_picture(BytesIO(data))
        except Exception as exc:
            _log.warning("inline image %r could not be embedded: %s", node.src, exc)
            self._add_inline_image_placeholder(host, node.alt)

    def _add_inline_image_placeholder(self, host: Any, alt: str) -> None:
        """Render a compact icon plus the same localized label as block images."""

        run = self._make_run(
            host,
            "",
            bold=False,
            italic=False,
            strike=False,
            underline=False,
        )
        if run is None:
            return
        run.add_picture(BytesIO(placeholder_png()), width=Pt(18))
        self._add_run(
            host,
            f"[{placeholder_copy(alt)}]",
            bold=False,
            italic=False,
            strike=False,
            underline=False,
        )

    # ---- render -------------------------------------------------------------

    def render(
        self,
        previous_rendered: RenderedInfo | None,
        layout_state: LayoutState,
    ) -> Generator[RenderedInfo | SubRenderable]:
        remaining_space = layout_state.remaining_page_height

        if self.page_break_before:
            layout_state.add_height(layout_state.remaining_page_height)

        previous_docx_paragraph: DocxParagraph | None = None
        if previous_rendered is not None and isinstance(
            previous_rendered.docx_element, DocxParagraph
        ):
            previous_docx_paragraph = previous_rendered.docx_element

        height_data = ParagraphSizer(
            self._docx_paragraph,
            previous_docx_paragraph,
            layout_state.max_width,
        ).calculate_height()

        if layout_state.current_page_height == 0 and layout_state.page > 1:
            height_data.before = Length(0)

        fitting_lines = 0
        for lines in range(1, height_data.lines + 1):
            partial = (
                height_data.before
                + ((lines - 1) * height_data.line_spacing + 1) * height_data.line_height
            )
            if partial > layout_state.remaining_page_height:
                break
            fitting_lines += 1

        height: float
        if fitting_lines == height_data.lines:
            height = float(min(height_data.full, layout_state.remaining_page_height))
        elif fitting_lines <= 1 or (
            height_data.lines - fitting_lines == 1 and height_data.lines == 3
        ):
            height = float(layout_state.remaining_page_height + height_data.full)
        elif height_data.lines - fitting_lines == 1:
            height = float(
                layout_state.remaining_page_height
                + height_data.before
                + height_data.line_height * height_data.line_spacing * 2
                + height_data.after
            )
        else:
            height = float(
                layout_state.remaining_page_height
                + height_data.before
                + height_data.line_height
                * height_data.line_spacing
                * (height_data.lines - fitting_lines)
                + height_data.after
            )

        if self.page_break_before:
            height += float(remaining_space)

        final_height = Length(int(height))
        yield RenderedInfo(self._docx_paragraph, final_height)
        layout_state.add_height(final_height)
